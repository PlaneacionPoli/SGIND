"""Asociación PRELIMINAR de línea y objetivo (PDI vigente) para los indicadores de la hoja
PDI_2026_2030 que aún no la tienen.

El criterio sale del contenido del indicador (nombre, descripción, proceso) frente a los
objetivos oficiales del PDI. Es una propuesta para que Planeación la revise, no una decisión:
  - Solo llena Linea y Objetivo cuando están VACÍOS; nunca pisa lo ya diligenciado.
  - NO toca PDI (1/0) ni Meta: la meta estratégica la define Planeación.
  - Marca la fila en Observaciones («Preliminar: línea y objetivo propuestos por contenido»).

Uso: backend/.venv312/Scripts/python.exe scripts/asociar_preliminar_pdi.py [--dry-run]
"""

from __future__ import annotations

import argparse
import os
import shutil
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR / "scripts"))
sys.path.insert(0, str(BASE_DIR / "backend"))
os.environ.setdefault("SGIND_DATA_PATH", str(BASE_DIR / "data"))

import agregar_hojas_marco_catalogo as hojas  # noqa: E402
from app.domain.marcos import get_marco  # noqa: E402
from app.domain.taxonomia import load_taxonomia  # noqa: E402

VERSION = "PDI-2026-2030"
NOTA = "Preliminar: línea y objetivo propuestos por contenido; revisar"

# Objetivo (código oficial) → Ids de indicador. Criterio resumido por grupo.
ASIGNACION: dict[str, str] = {
    # L1 · Calidad e Innovación Educativa
    "L1-OI": "473 474 476 477 478 479 480 522 PRY-28 PRY-43 PRY-52",  # modelo educativo / currículo / resultados de aprendizaje
    "L1-OII": (  # enseñanza y aprendizaje con tecnología; formación docente
        "268 269 270 271 302 487 401 402 403 544 626 627 397 453 455 448 282 283 285 288 460 576 616 617"
    ),
    "L1-OIII": (  # mejoramiento continuo, acreditación, calidad de procesos, investigación
        "68 69 122 172 436 481 485 488 290 526 262 263 265 266 409 143 144 145 146 635 255 256 543 "
        "239 240 241 509 PRY-32 251 306 254 373 390 414 415 416 417 418 420 469 470 471 528 "
        "273 326 327 337 338 475 636 277 278 286 358 336 197 300 301 529 "
        "510 511 513 514 518 538 PRY-54 PRY-26 527 558 559 563 565 566 573 574 578 603 604 605 606 607 608 595 612"
    ),
    # L2 · Experiencia Centrada en las Personas
    "L2-OI": (  # comunidad Poli: bienestar, ADN, talento humano, SST, cultura, salud
        "73 74 167 199 200 202 332 334 344 346 556 PRY-30 "
        "84 85 86 87 88 89 90 91 92 94 95 100 104 106 124 125 126 127 128 272 431 432 484 "
        "292 294 303 304 447 109 426 456 457 486 PRY-45 597 599 600 601 602 615"
    ),
    "L2-OII": (  # experiencias, satisfacción, servicio, datos y ecosistema digital
        "111 112 113 114 115 116 117 118 133 150 151 159 437 490 491 492 493 494 495 496 497 498 499 500 501 "
        "502 503 504 505 506 507 508 552 553 554 619 622 632 633 634 275 359 466 467 468 547 625 "
        "245 252 389 246 444 445 539 260 261 193 537 620 623 362 363 366 375 439 440 "
        "540 624 637 638 639 640 641 515 530 531 "
        "549 560 564 567 575 579 580 583 584 585 587 588 593 594 598 609 610 613 614 PRY-31 PRY-42 PRY-46 PRY-47 PRY-48 PRY-49 PRY-50 PRY-51"
    ),
    # L3 · Expansión con Compromiso Social
    "L3-OI": (  # relacionamiento, posicionamiento y crecimiento: matrícula, mercadeo, comunicaciones
        "379 381 382 383 384 274 422 423 424 524 PRY-53 394 395 413 451 452 519 392 393 "
        "630 631 318 541 542 546 PRY-40 562 581 589 590 591 592 611"
    ),
    "L3-OII": "532 533 534 535 536 PRY-44",  # presencia en América / internacionalización
    "L3-OIII": (  # portafolio de servicios educativos
        "385 386 368 372 555 621 324 408 PRY-12 PRY-13 PRY-17"
    ),
    # L4 · Desarrollo Sostenible
    "L4-OI": (  # gestión sostenible: ambiental, RSU, inclusión, emprendimiento
        "208 213 214 215 216 217 218 219 220 221 222 223 224 225 226 227 228 230 231 232 233 234 235 236 238 "
        "244 427 428 429 433 434 435 628 629 361 364 367 369 398 399 400 550 551 557 568 569 570 571 572 586"
    ),
    "L4-OII": (  # recursos financieros
        "204 205 206 207 355 356 357 377 430 347 348 351 352 341 121 371 561 577 596"
    ),
}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args(argv)

    tax = load_taxonomia(get_marco(VERSION).taxonomia)
    por_codigo = {o.id: (ln.nombre, o.nombre) for ln in tax.lineas for o in ln.objetivos}
    destino: dict[str, str] = {}
    for cod, ids in ASIGNACION.items():
        if cod not in por_codigo:
            raise SystemExit(f"Objetivo inexistente en la taxonomía: {cod}")
        for i in ids.split():
            if i in destino:
                raise SystemExit(f"Indicador {i} asignado dos veces ({destino[i]} y {cod})")
            destino[i] = cod

    wb = openpyxl.load_workbook(hojas.CATALOGO_FILE)
    ws = wb[hojas.hoja_pdi(get_marco(VERSION))]
    cols = [c.value for c in ws[1]]
    i_id, i_l, i_o, i_obs = (cols.index(c) + 1 for c in ("Id", "Linea", "Objetivo", "Observaciones"))

    llenadas: Counter = Counter()
    sin_regla: list[str] = []
    for fila in range(2, ws.max_row + 1):
        ids = str(ws.cell(row=fila, column=i_id).value or "").strip()
        if not ids:
            continue
        vacia_l = ws.cell(row=fila, column=i_l).value in (None, "")
        vacia_o = ws.cell(row=fila, column=i_o).value in (None, "")
        if not (vacia_l or vacia_o):
            continue
        cod = destino.get(ids)
        if cod is None:
            sin_regla.append(ids)
            continue
        linea, objetivo = por_codigo[cod]
        if vacia_l:
            ws.cell(row=fila, column=i_l, value=linea)
        if vacia_o:
            ws.cell(row=fila, column=i_o, value=objetivo)
        obs = ws.cell(row=fila, column=i_obs)
        obs.value = f"{obs.value}. {NOTA}" if obs.value else NOTA
        llenadas[cod] += 1

    print(f"Completados: {sum(llenadas.values())}")
    for cod in sorted(llenadas):
        print(f"  {cod:8} {llenadas[cod]:3}  {por_codigo[cod][0]}")
    if sin_regla:
        print(f"Sin criterio ({len(sin_regla)}): {', '.join(sin_regla)}")
    sobran = sorted(set(destino) - {str(ws.cell(row=r, column=i_id).value).strip() for r in range(2, ws.max_row + 1)})
    if sobran:
        print(f"Ids de la tabla que no están en la hoja: {', '.join(sobran)}")
    if args.dry_run:
        print("--dry-run: no se escribió nada.")
        return 0

    hojas.BACKUP_DIR = BASE_DIR / "data" / "raw" / "_backups_fusion"
    hojas.BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(hojas.CATALOGO_FILE, hojas.BACKUP_DIR / f"Catalogo de Indicadores.{datetime.now():%Y%m%d_%H%M%S}.pre_preliminar.xlsx")
    try:
        wb.save(hojas.CATALOGO_FILE)
    except PermissionError:
        print(f"[ERROR] Cierra {hojas.CATALOGO_FILE.name} en Excel y reintenta.")
        return 3
    print(f"OK {hojas.CATALOGO_FILE.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
