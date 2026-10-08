"""Importa la taxonomía OFICIAL de un PDI (líneas, objetivos y metas) desde su Excel y
actualiza todo lo que depende de ella.

Fuente: data/raw/PDI_<ciclo>_Lineas_Objetivos_Metas.xlsx, hoja con las columnas
  Cód. Línea | Línea Estratégica | Cód. Objetivo | Objetivo Estratégico | Cód. Meta | Meta Estratégica

Qué hace:
  1. Valida la estructura (líneas, objetivos y metas únicos, códigos coherentes).
  2. Reescribe backend/app/data/taxonomia/<version>.json con los textos oficiales y los
     códigos del Excel como ids (L1, L1-OI, L1-OI-M1).
  3. Migra lo ya diligenciado en la hoja PDI_<ciclo> del catálogo: los objetivos y metas
     escogidos con el texto anterior pasan al texto oficial (se emparejan por posición
     dentro de la misma línea/objetivo, porque la estructura es la misma). Lo que no se
     pueda emparejar se informa y NO se modifica.
  4. Regenera Taxonomia_Marco y Listas_PDI_<ciclo> (desplegables en cascada) y reaplica
     las validaciones de la hoja del PDI.

Uso:
    backend/.venv312/Scripts/python.exe scripts/importar_taxonomia_pdi.py
    ... --version PDI-2026-2030 --excel RUTA --dry-run
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR / "scripts"))
sys.path.insert(0, str(BASE_DIR / "backend"))
os.environ.setdefault("SGIND_DATA_PATH", str(BASE_DIR / "data"))

import agregar_hojas_marco_catalogo as hojas  # noqa: E402
from app.domain.marcos import get_marco  # noqa: E402
from app.domain.taxonomia import TAXONOMIA_DIR, load_taxonomia, norm_texto  # noqa: E402

ENCABEZADOS = ("cod linea", "linea estrategica", "cod objetivo", "objetivo estrategico", "cod meta", "meta estrategica")
_ROMANOS = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}


def leer_excel(path: Path) -> list[tuple[str, str, str, str, str, str]]:
    """Filas (cod_l, linea, cod_o, objetivo, cod_m, meta) de la primera hoja con los encabezados."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb:
        filas = list(ws.iter_rows(values_only=True))
        for i, fila in enumerate(filas):
            if tuple(norm_texto(c) for c in fila[:6]) == ENCABEZADOS:
                datos = []
                for r in filas[i + 1:]:
                    if r[0] is None or not re.match(r"^L\d+$", str(r[0]).strip()):
                        break
                    datos.append(tuple(str(c).strip() for c in r[:6]))
                return datos
    raise SystemExit(f"No se encontró la tabla con encabezados {ENCABEZADOS} en {path.name}")


def construir_json(version_id: str, filas, fuente: str) -> dict:
    lineas: dict[str, dict] = {}
    for cod_l, linea, cod_o, objetivo, cod_m, meta in filas:
        ln = lineas.setdefault(cod_l, {"id": cod_l, "orden": len(lineas) + 1, "nombre": linea, "objetivos": {}})
        if ln["nombre"] != linea:
            raise SystemExit(f"{cod_l}: dos nombres de línea distintos ({ln['nombre']!r} / {linea!r})")
        ob = ln["objetivos"].setdefault(cod_o, {"id": cod_o, "numero": len(ln["objetivos"]) + 1, "nombre": objetivo, "metas": {}})
        if ob["nombre"] != objetivo:
            raise SystemExit(f"{cod_o}: dos textos de objetivo distintos")
        if not cod_o.startswith(cod_l + "-") or not cod_m.startswith(cod_o + "-"):
            raise SystemExit(f"Códigos incoherentes: {cod_l} / {cod_o} / {cod_m}")
        if cod_m in ob["metas"]:
            raise SystemExit(f"Meta duplicada: {cod_m}")
        ob["metas"][cod_m] = {"id": cod_m, "nombre": meta}
    return {
        "version_id": version_id,
        "tipo": "PDI",
        "fuente": fuente,
        "lineas": [
            {**ln, "objetivos": [{**ob, "metas": list(ob["metas"].values())} for ob in ln["objetivos"].values()]}
            for ln in lineas.values()
        ],
    }


def equivalencias(viejo, nuevo) -> dict[str, dict[str, str]]:
    """{'objetivo': {texto viejo normalizado: texto nuevo}, 'meta': {...}} emparejando por posición."""
    eq: dict[str, dict[str, str]] = {"objetivo": {}, "meta": {}}
    if viejo is None or len(viejo.lineas) != len(nuevo["lineas"]):
        return eq
    for lv, ln in zip(viejo.lineas, nuevo["lineas"], strict=True):
        if len(lv.objetivos) != len(ln["objetivos"]):
            continue
        for ov, on in zip(lv.objetivos, ln["objetivos"], strict=True):
            eq["objetivo"][norm_texto(ov.nombre)] = on["nombre"]
            if len(ov.metas) == len(on["metas"]):
                for mv, mn in zip(ov.metas, on["metas"], strict=True):
                    eq["meta"][norm_texto(mv.nombre)] = mn["nombre"]
    return eq


def migrar_hoja(wb, marco, eq, nuevo_json) -> tuple[int, list[str]]:
    """Reescribe Objetivo/Meta de la hoja del PDI al texto oficial. Devuelve (cambios, sin_emparejar)."""
    nombre = hojas.hoja_pdi(marco)
    if nombre not in wb.sheetnames:
        return 0, []
    ws = wb[nombre]
    cols = [c.value for c in ws[1]]
    oficiales = {
        "objetivo": {norm_texto(o["nombre"]) for ln in nuevo_json["lineas"] for o in ln["objetivos"]},
        "meta": {norm_texto(m["nombre"]) for ln in nuevo_json["lineas"] for o in ln["objetivos"] for m in o["metas"]},
    }
    cambios, sin = 0, []
    for fila in range(2, ws.max_row + 1):
        for campo, col in (("objetivo", "Objetivo"), ("meta", "Meta")):
            celda = ws.cell(row=fila, column=cols.index(col) + 1)
            if celda.value in (None, ""):
                continue
            clave = norm_texto(celda.value)
            if clave in oficiales[campo]:
                if celda.value != next(  # mismo texto salvo puntuación/tildes: se deja el oficial
                    t for t in _textos(nuevo_json, campo) if norm_texto(t) == clave
                ):
                    celda.value = next(t for t in _textos(nuevo_json, campo) if norm_texto(t) == clave)
                    cambios += 1
            elif clave in eq[campo]:
                celda.value = eq[campo][clave]
                cambios += 1
            else:
                sin.append(f"{ws.cell(row=fila, column=1).value}: {col} «{celda.value}»")
    return cambios, sin


def _textos(nuevo_json, campo):
    for ln in nuevo_json["lineas"]:
        for o in ln["objetivos"]:
            if campo == "objetivo":
                yield o["nombre"]
            else:
                for m in o["metas"]:
                    yield m["nombre"]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--version", default="PDI-2026-2030")
    ap.add_argument("--excel", type=Path, default=BASE_DIR / "data" / "raw" / "PDI_2026-2030_Lineas_Objetivos_Metas.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args(argv)

    marco = get_marco(args.version)
    filas = leer_excel(args.excel)
    nuevo = construir_json(args.version, filas, f"{args.excel.name} (textos oficiales del PDI)")
    n_l = len(nuevo["lineas"])
    n_o = sum(len(ln["objetivos"]) for ln in nuevo["lineas"])
    n_m = len(filas)
    print(f"Excel: {n_l} líneas, {n_o} objetivos, {n_m} metas")

    destino = TAXONOMIA_DIR / f"{marco.taxonomia}.json"
    viejo = load_taxonomia(marco.taxonomia) if destino.exists() else None
    eq = equivalencias(viejo, nuevo)
    print(f"Equivalencias texto anterior a oficial: {len(eq['objetivo'])} objetivos, {len(eq['meta'])} metas")

    if args.dry_run:
        print("--dry-run: no se escribió nada.")
        return 0

    destino.write_text(json.dumps(nuevo, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Escrito {destino.relative_to(BASE_DIR)}")
    load_taxonomia.cache_clear() if hasattr(load_taxonomia, "cache_clear") else None

    wb = openpyxl.load_workbook(hojas.CATALOGO_FILE)
    cambios, sin = migrar_hoja(wb, marco, eq, nuevo)
    print(f"{hojas.hoja_pdi(marco)}: {cambios} celdas migradas al texto oficial")
    for s in sin:
        print("  [SIN EMPAREJAR, no se modificó]", s)
    hojas.escribir_taxonomia(wb)
    listas = hojas.escribir_listas(wb, marco)
    if hojas.hoja_pdi(marco) in wb.sheetnames:
        hojas.aplicar_validaciones(wb[hojas.hoja_pdi(marco)], hojas.columnas(marco), marco, listas)
    try:
        wb.save(hojas.CATALOGO_FILE)
    except PermissionError:
        print(f"[ERROR] Cierra {hojas.CATALOGO_FILE.name} en Excel y reintenta (el JSON ya se actualizó).")
        return 3
    print(f"OK {hojas.CATALOGO_FILE.name}: Taxonomia_Marco, {listas['hoja']} y validaciones actualizadas")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
