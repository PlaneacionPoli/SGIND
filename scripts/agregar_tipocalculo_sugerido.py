"""
scripts/agregar_tipocalculo_sugerido.py
Agrega al catálogo maestro (hoja 'Catalogo Indicadores') el TipoCalculo SUGERIDO de los indicadores
activos que no tienen TipoCalculo. No toca la columna TipoCalculo: el ETL sigue leyendo solo esa;
quien valide una sugerencia la copia a TipoCalculo.

Columnas nuevas (al final de la hoja, para no desplazar las existentes):
  - TipoCalculo_Sugerido  (celda resaltada en amarillo)
  - TipoCalculo_Evidencia (de dónde sale la sugerencia y su confianza)

La sugerencia se obtiene por Id, con la hoja 'Ficha Tecnica Detalle' del mismo libro:
  - la Fórmula o Descripción menciona «acumulado / sumatoria / suma de»  → Acumulado (media)
  - la Fórmula o Descripción menciona «promedio»                         → Promedio  (media)
  - el nombre del indicador empieza por «Número de / Cantidad de»        → Acumulado (baja)
  - sin evidencia                                                         → Cierre    (baja; es lo que el ETL asume)

Es idempotente: vuelve a calcular solo las filas sin TipoCalculo y limpia la sugerencia de las que
ya lo tienen. openpyxl descarta las listas desplegables «extensión» (x14: Línea→Objetivo→Meta de las
hojas PDI); este script las reinyecta tras guardar.

Uso:  python scripts/agregar_tipocalculo_sugerido.py [--dry-run]
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent.parent
CATALOGO = ROOT / "data" / "raw" / "Catalogo de Indicadores.xlsx"
HOJA = "Catalogo Indicadores"
HOJA_FICHA = "Ficha Tecnica Detalle"
COL_SUGERIDO = "TipoCalculo_Sugerido"
COL_EVIDENCIA = "TipoCalculo_Evidencia"
FILL_SUGERIDO = PatternFill("solid", fgColor="FFF2CC")

_ACUMULADO = re.compile(r"acumulad|sumatoria|suma de", re.I)
_PROMEDIO = re.compile(r"promedio", re.I)
_CONTEO = re.compile(r"^\s*(n[uú]mero|cantidad) de\b", re.I)


def _txt(v: object) -> str:
    return "" if v is None else str(v).strip()


def sugerir(nombre: str, texto_ficha: str) -> tuple[str, str]:
    """(tipo sugerido, evidencia con confianza) para un indicador sin TipoCalculo."""
    if _ACUMULADO.search(texto_ficha):
        return "Acumulado", "Ficha técnica (Id): menciona acumulado/suma · confianza media"
    if _PROMEDIO.search(texto_ficha):
        return "Promedio", "Ficha técnica (Id): menciona promedio · confianza media"
    if _CONTEO.search(nombre):
        return "Acumulado", "Nombre: conteo (Número/Cantidad de) · confianza baja"
    return "Cierre", "Sin evidencia en la ficha: valor por defecto del ETL · confianza baja"


def _indice_ficha(wb) -> dict[str, str]:
    """Id → texto de Fórmula + Descripción en la ficha técnica."""
    if HOJA_FICHA not in wb.sheetnames:
        return {}
    ws = wb[HOJA_FICHA]
    enc = [_txt(c.value) for c in ws[1]]
    i_id = enc.index("Id")
    i_txt = [i for i, c in enumerate(enc) if c == "Formula" or c.startswith("Descripci")]
    out: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[i_id] is None:
            continue
        ids = _txt(row[i_id]).removesuffix(".0")
        out.setdefault(ids, " | ".join(_txt(row[i]) for i in i_txt if i < len(row) and row[i]))
    return out


def agregar_sugeridos(wb) -> dict[str, int]:
    ws = wb[HOJA]
    enc = [_txt(c.value) for c in ws[1]]
    i_id, i_nom, i_est, i_tc = (enc.index(c) for c in ("Id", "Indicador", "Estado", "TipoCalculo"))
    ficha = _indice_ficha(wb)

    def _col(nombre: str) -> int:
        if nombre in enc:
            return enc.index(nombre) + 1
        ws.cell(1, len(enc) + 1).value = nombre
        ws.cell(1, len(enc) + 1).font = ws.cell(1, 1).font.copy()
        ws.cell(1, len(enc) + 1).fill = ws.cell(1, 1).fill.copy()
        enc.append(nombre)
        return len(enc)

    c_sug, c_evi = _col(COL_SUGERIDO), _col(COL_EVIDENCIA)
    cuenta = {"sugeridos": 0, "limpiados": 0}
    for row in ws.iter_rows(min_row=2):
        if row[i_id].value is None:
            continue
        sug, evi = ws.cell(row[0].row, c_sug), ws.cell(row[0].row, c_evi)
        if _txt(row[i_tc].value):  # ya tiene TipoCalculo: no necesita sugerencia
            if sug.value is not None or evi.value is not None:
                sug.value = evi.value = None
                sug.fill = PatternFill(fill_type=None)
                cuenta["limpiados"] += 1
            continue
        if _txt(row[i_est].value).lower() != "activo":
            continue
        ids = _txt(row[i_id].value).removesuffix(".0")
        sug.value, evi.value = sugerir(_txt(row[i_nom].value), ficha.get(ids, ""))
        sug.fill = FILL_SUGERIDO
        cuenta["sugeridos"] += 1
    ws.column_dimensions[openpyxl.utils.get_column_letter(c_sug)].width = 22
    ws.column_dimensions[openpyxl.utils.get_column_letter(c_evi)].width = 60
    return cuenta


# ── Preservar las listas desplegables «extensión» (x14) que openpyxl descarta ──────────────────

_EXTLST = re.compile(r"<extLst>.*?</extLst>", re.S)


def _rutas_hojas(z: zipfile.ZipFile) -> dict[str, str]:
    """ruta interna de cada hoja → nombre de la hoja."""
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    destino = {}
    for rel in re.findall(r"<Relationship [^>]*>", rels):
        destino[re.search(r'Id="([^"]+)"', rel).group(1)] = re.search(r'Target="([^"]+)"', rel).group(1)
    return {
        "xl/" + destino[rid].lstrip("/").removeprefix("xl/"): nombre
        for nombre, rid in re.findall(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="([^"]+)"', wbx)
    }


def _extlst_por_hoja(xlsx: Path) -> dict[str, str]:
    """nombre de hoja → bloque <extLst> con x14:dataValidations (si lo tiene)."""
    out: dict[str, str] = {}
    with zipfile.ZipFile(xlsx) as z:
        for ruta, nombre in _rutas_hojas(z).items():
            m = _EXTLST.search(z.read(ruta).decode("utf-8"))
            if m and "x14:dataValidation" in m.group(0):
                out[nombre] = m.group(0)
    return out


def _reinyectar(xlsx: Path, bloques: dict[str, str]) -> int:
    """Inserta cada <extLst> en la hoja del mismo nombre del libro recién guardado."""
    if not bloques:
        return 0
    n = 0
    tmp = xlsx.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(xlsx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        rutas = _rutas_hojas(zin)
        for item in zin.infolist():
            data = zin.read(item.filename)
            nombre = rutas.get(item.filename)
            if nombre in bloques:
                xml = data.decode("utf-8")
                bloque = re.sub(r'\s+xr:uid="[^"]*"', "", bloques[nombre])  # xr no se declara en el libro nuevo
                if "<extLst>" not in xml:
                    data = xml.replace("</worksheet>", bloque + "</worksheet>").encode("utf-8")
                    n += 1
            zout.writestr(item, data)
    shutil.move(tmp, xlsx)
    return n


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="calcula y muestra, sin guardar")
    ap.add_argument("--catalogo", type=Path, default=CATALOGO, help="libro a actualizar (por defecto el maestro)")
    args = ap.parse_args()
    destino_final: Path = args.catalogo
    if not destino_final.exists():
        sys.exit(f"No existe {destino_final}")

    bloques = _extlst_por_hoja(destino_final)
    wb = openpyxl.load_workbook(destino_final)
    cuenta = agregar_sugeridos(wb)
    print(f"Sugeridos: {cuenta['sugeridos']} | limpiados (ya tienen TipoCalculo): {cuenta['limpiados']}")
    if args.dry_run:
        return
    with tempfile.TemporaryDirectory(prefix="cat_sug_") as td:
        salida = Path(td) / destino_final.name
        wb.save(salida)
        n = _reinyectar(salida, bloques)
        shutil.copy2(salida, destino_final)
    print(f"Guardado {destino_final.name}; listas desplegables x14 reinyectadas en {n} hoja(s)")


if __name__ == "__main__":
    main()
