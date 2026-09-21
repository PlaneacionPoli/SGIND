"""
scripts/cna_extraction/diagnostics.py

Construye el reporte de validación de la Fase 1 (solo lectura): recorre el
catálogo, resuelve cada hoja y corre el detector de estructura, agregando
conteos y un listado de casos que requieren revisión manual antes de
implementar la Fase 2 (normalización + escritura incremental).
"""

from __future__ import annotations

import json
import logging
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger(__name__)

from scripts.cna_extraction.catalog import load_catalog
from scripts.cna_extraction.sheet_resolver import resolve_sheets
from scripts.cna_extraction.structure_detector import detect_structure
from scripts.etl.audit import AuditTrail

ATYPICAL_KINDS = {"unresolved", "event_log_atypical"}


def run_diagnostics(xlsx_path: Path) -> dict[str, Any]:
    """Ejecuta el diagnóstico completo (catálogo -> resolver -> detector)
    contra `xlsx_path` y devuelve un dict serializable con los conteos y la
    lista de `problems` a revisar. No escribe nada en data/output/."""
    xlsx_path = Path(xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    try:
        catalog = load_catalog(xlsx_path)
        resolved, unreferenced_sheets = resolve_sheets(catalog, wb)
        total_sheets = len(wb.sheetnames)

        kind_counts: Counter[str] = Counter()
        period_histogram: Counter[str] = Counter()
        subindicador_depth_histogram: Counter[int] = Counter()
        problems: list[dict[str, Any]] = []
        sheets_matched = 0
        sheets_procesados = 0

        for item in resolved:
            if item.sheet_name is None:
                problems.append(
                    {
                        "numero": item.catalog_record.numero if item.catalog_record else None,
                        "sheet_name": None,
                        "catalog_nombre": item.catalog_record.nombre if item.catalog_record else None,
                        "kind": "sin_hoja",
                        "reason": "Registro de catálogo sin hoja correspondiente en el workbook.",
                    }
                )
                continue

            if item.catalog_record is None:
                problems.append(
                    {
                        "numero": None,
                        "sheet_name": item.sheet_name,
                        "catalog_nombre": None,
                        "kind": "sin_catalogo",
                        "reason": "Hoja del workbook sin registro correspondiente en Índice Tablas.",
                    }
                )
                continue

            sheets_matched += 1
            ws = wb[item.sheet_name]
            raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
            numero = item.catalog_record.numero
            structure = detect_structure(raw_rows, numero)
            sheets_procesados += 1
            kind_counts[structure.kind] += 1

            for r in structure.rows:
                if r.get("period"):
                    period_histogram[r["period"]] += 1

            depths = {len(r["category_path"]) for r in structure.rows}
            for d in depths:
                subindicador_depth_histogram[d] += 1

            if structure.kind in ATYPICAL_KINDS:
                problems.append(
                    {
                        "numero": numero,
                        "sheet_name": item.sheet_name,
                        "catalog_nombre": item.catalog_record.nombre,
                        "kind": structure.kind,
                        "reason": "; ".join(structure.issues) or "Estructura no interpretable automáticamente.",
                    }
                )
            elif structure.has_explicit_total is False and any(len(r["category_path"]) >= 1 for r in structure.rows):
                problems.append(
                    {
                        "numero": numero,
                        "sheet_name": item.sheet_name,
                        "catalog_nombre": item.catalog_record.nombre,
                        "kind": "sin_total_explicito",
                        "reason": (
                            "Tiene subdivisiones pero no trae fila/columna Total explícita; "
                            "la Fase 2 debe calcular el valor del indicador principal como "
                            "suma de las subdivisiones por periodo."
                        ),
                    }
                )

        report: dict[str, Any] = {
            "generado": datetime.now().isoformat(),
            "archivo_fuente": str(xlsx_path),
            "total_sheets_workbook": total_sheets,
            "total_catalog_records": len(catalog),
            "total_sheets_datos": len({r.sheet_name for r in resolved if r.sheet_name}),
            "sheets_matched": sheets_matched,
            "sheets_procesados": sheets_procesados,
            "unreferenced_sheets": unreferenced_sheets,
            "sheets_by_kind": dict(kind_counts),
            "periods_found_histogram": dict(sorted(period_histogram.items())),
            "subindicador_depth_histogram": dict(sorted(subindicador_depth_histogram.items())),
            "problems": problems,
        }
        return report
    finally:
        wb.close()


def write_report(report: dict[str, Any], out_dir: Path = Path("artifacts")) -> tuple[Path, Path]:
    """Escribe el reporte en artifacts/cna_extraction_reporte_<timestamp>.{json,md},
    siguiendo la convención ya usada en artifacts/reporte_YYYYMMDD.*, y
    registra la corrida en la bitácora de auditoría compartida (solo
    lectura: nada se escribe en data/output/ en esta fase)."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = out_dir / f"cna_extraction_reporte_{stamp}.json"
    md_path = out_dir / f"cna_extraction_reporte_{stamp}.md"

    json_path.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    lines = [
        "# Reporte de diagnóstico — extracción CNA (Fase 1)",
        "",
        f"- Archivo fuente: `{report['archivo_fuente']}`",
        f"- Generado: {report['generado']}",
        f"- Hojas totales del workbook: {report['total_sheets_workbook']}",
        f"- Registros de catálogo (Índice Tablas): {report['total_catalog_records']}",
        f"- Hojas de datos en workbook: {report['total_sheets_datos']}",
        f"- Hojas emparejadas con catálogo: {report['sheets_matched']}",
        f"- Hojas procesadas por el detector: {report['sheets_procesados']}",
        "",
        "## Hojas por tipo de estructura detectada",
        "",
    ]
    for kind, count in sorted(report["sheets_by_kind"].items()):
        lines.append(f"- {kind}: {count}")

    lines += ["", "## Hojas del workbook sin correspondencia en el catálogo", ""]
    if report["unreferenced_sheets"]:
        for s in report["unreferenced_sheets"]:
            lines.append(f"- {s}")
    else:
        lines.append("(ninguna)")

    lines += [
        "",
        "## Casos a revisar manualmente",
        "",
        "| Número | Hoja | Indicador | Tipo | Motivo |",
        "|---|---|---|---|---|",
    ]
    for p in report["problems"]:
        lines.append(f"| {p['numero']} | {p['sheet_name']} | {p['catalog_nombre']} | {p['kind']} | {p['reason']} |")

    md_path.write_text("\n".join(lines), encoding="utf-8")

    try:
        trail = AuditTrail()
        trail.registrar_ejecucion(
            evento="diagnostico_cna_extraction",
            detalles={
                "archivo_fuente": report["archivo_fuente"],
                "sheets_procesados": report["sheets_procesados"],
                "problemas_detectados": len(report["problems"]),
                "reporte_json": str(json_path),
            },
            exitoso=True,
        )
    except Exception as exc:
        # La auditoría no debe bloquear la generación del reporte de diagnóstico
        # (ej. archivo de auditoría bloqueado por OneDrive/otro proceso).
        logger.warning("No se pudo registrar la corrida en la bitácora de auditoría: %s", exc)

    return json_path, md_path
