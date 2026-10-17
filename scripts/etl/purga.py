"""
scripts/etl/purga.py
Limpieza y deduplicación de hojas Excel existentes antes de escribir nuevos datos.
"""
from __future__ import annotations

import logging
from collections import defaultdict
from datetime import date
from typing import Dict, List, Optional, Set, Tuple

import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill

from .config import AÑO_CIERRE_ACTUAL
from .escritura import _ejec_score, get_last_data_row
from .extraccion import (
    _ejec_corrected_from_row, _meta_corrected_from_row,
    _usa_variables_canonico, _IDS_META_FIJA, _IDS_SUMA_VARIABLES_SERIES,
)
from .formulas_excel import _build_col_map
from .normalizacion import _id_str, make_llave
from .periodos import ultimo_dia_mes

logger = logging.getLogger(__name__)



# ── Purga de filas inválidas ──────────────────────────────────────

def purgar_filas_invalidas(
    ws,
    nombre: str = "hoja",
    kawak_validos: Optional[Set[Tuple[str, int]]] = None,
) -> int:
    """
    Elimina filas donde:
      - La fecha es futura (año > AÑO_CIERRE_ACTUAL)
      - El campo Año contiene texto inválido
      - El par (Id, año) no existe en el catálogo Kawak (si kawak_validos != None),
        pero SOLO si la fila no tiene dato real: Meta y Ejecución ambas vacías/0.
        El catálogo Kawak solo cubre años con reporte por API; años anteriores con
        valor cargado manualmente en la fuente original (Resultados_Consolidados_
        Fuente.xlsx) son datos reales y no deben destruirse (feedback 2026-08-13:
        purga borró silenciosamente 2022 de indicadores 323/276/423 y otros que
        empezaron a reportarse por API solo desde 2023).
    """
    cm = _build_col_map(ws)
    idx_id    = cm.get("Id",    1) - 1
    idx_fecha = cm.get("Fecha", 6) - 1
    idx_anio  = cm.get("Anio",  7) - 1
    idx_meta  = cm.get("Meta")
    idx_ejec  = cm.get("Ejecucion")
    idx_meta  = idx_meta - 1 if idx_meta is not None else None
    idx_ejec  = idx_ejec - 1 if idx_ejec is not None else None

    def _vacio(v) -> bool:
        if v is None:
            return True
        try:
            return float(v) == 0.0
        except (TypeError, ValueError):
            return str(v).strip() in ("", "nan", "None")

    filas_a_borrar = []
    n_kawak = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        fecha_raw = row[idx_fecha].value if len(row) > idx_fecha else None
        año_fila  = None
        try:
            fecha = pd.to_datetime(fecha_raw)
            año_fila = fecha.year
            if fecha.year > AÑO_CIERRE_ACTUAL:
                filas_a_borrar.append(row[0].row)
                continue
        except Exception:
            pass
        anio_val = row[idx_anio].value if len(row) > idx_anio else None
        if anio_val is not None:
            if isinstance(anio_val, str) and anio_val.startswith("="):
                pass
            else:
                try:
                    año_fila = int(float(anio_val))
                except (TypeError, ValueError):
                    filas_a_borrar.append(row[0].row)
                    continue
        if kawak_validos is not None and año_fila is not None:
            id_val = row[idx_id].value if len(row) > idx_id else None
            id_s   = _id_str(id_val) if id_val is not None else None
            if id_s and (id_s, año_fila) not in kawak_validos:
                meta_val = row[idx_meta].value if idx_meta is not None and len(row) > idx_meta else None
                ejec_val = row[idx_ejec].value if idx_ejec is not None and len(row) > idx_ejec else None
                if _vacio(meta_val) and _vacio(ejec_val):
                    filas_a_borrar.append(row[0].row)
                    n_kawak += 1

    for r_idx in sorted(set(filas_a_borrar), reverse=True):
        ws.delete_rows(r_idx)
    total = len(set(filas_a_borrar))
    if total:
        logger.info(
            f"  [{nombre}] {total} filas eliminadas "
            f"({n_kawak} por no estar en Indicadores Kawak)."
        )
    return total


# ── Purga de filas anteriores a la Fecha Desde de la ficha técnica ─

def purgar_filas_antes_fecha_desde(
    ws, fecha_desde_map: Dict, nombre: str = "hoja",
) -> int:
    """
    Elimina filas de fecha anterior a la 'Fecha Desde' oficial del
    indicador (ficha técnica), pero SOLO si no tienen dato real: Meta y
    Ejecución ambas en 0 o vacías. Si hay un valor real reportado antes
    de esa fecha, la fila se conserva (feedback: no destruir datos).
    """
    cm = _build_col_map(ws)
    idx_id    = cm.get("Id",    1) - 1
    idx_fecha = cm.get("Fecha", 6) - 1
    idx_meta  = cm.get("Meta")
    idx_ejec  = cm.get("Ejecucion")
    if idx_meta is None or idx_ejec is None:
        return 0
    idx_meta -= 1
    idx_ejec -= 1

    def _vacio(v) -> bool:
        if v is None:
            return True
        try:
            return float(v) == 0.0
        except (TypeError, ValueError):
            return str(v).strip() in ("", "nan", "None")

    filas_a_borrar = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        id_s = _id_str(row[idx_id].value) if len(row) > idx_id else None
        fecha_desde = fecha_desde_map.get(id_s) if id_s else None
        if fecha_desde is None:
            continue
        try:
            fecha = pd.to_datetime(row[idx_fecha].value)
        except Exception:
            continue
        if pd.isna(fecha) or fecha >= fecha_desde:
            continue
        meta_val = row[idx_meta].value if len(row) > idx_meta else None
        ejec_val = row[idx_ejec].value if len(row) > idx_ejec else None
        if _vacio(meta_val) and _vacio(ejec_val):
            filas_a_borrar.append(row[0].row)

    for r_idx in sorted(set(filas_a_borrar), reverse=True):
        ws.delete_rows(r_idx)
    total = len(set(filas_a_borrar))
    if total:
        logger.info(
            f"  [{nombre}] {total} filas eliminadas por ser anteriores a "
            "Fecha Desde (ficha técnica) sin dato real."
        )
    else:
        logger.info(f"  [{nombre}] Sin filas anteriores a Fecha Desde para purgar.")
    return total


# ── Limpiar Consolidado Cierres ───────────────────────────────────

def limpiar_cierres_existentes(ws) -> int:
    """
    Años cerrados (anteriores al año en curso): conserva solo el registro de diciembre
    (o el último corte si no hay diciembre). Año en curso: conserva todos los cortes;
    luego `reemplazar_cierres_parciales` deja uno solo por Id.
    """
    curso = anio_en_curso()
    filas = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        fecha_raw = row[5].value
        try:    fecha = pd.to_datetime(fecha_raw)
        except: fecha = None
        filas.append({
            "row_idx": row[0].row,
            "Id":      row[0].value,
            "fecha":   fecha,
            "mes":     fecha.month if fecha else None,
            "año":     fecha.year  if fecha else None,
        })

    if not filas:
        return 0

    grupos: Dict = defaultdict(list)
    for f in filas:
        if f["año"] is None:
            continue
        grupos[(str(f["Id"]), f["año"])].append(f)

    filas_a_conservar: Set[int] = set()
    for (id_val, año), grupo in grupos.items():
        if año >= curso:
            for f in grupo:
                filas_a_conservar.add(f["row_idx"])
        else:
            dic = [f for f in grupo if f["mes"] == 12]
            keep = sorted(dic if dic else grupo, key=lambda f: f["fecha"])[-1]
            filas_a_conservar.add(keep["row_idx"])

    for f in filas:
        if f["año"] is None:
            filas_a_conservar.add(f["row_idx"])

    filas_a_borrar = sorted(
        [f["row_idx"] for f in filas if f["row_idx"] not in filas_a_conservar],
        reverse=True,
    )
    for r_idx in filas_a_borrar:
        ws.delete_rows(r_idx)

    logger.info(f"  limpiar_cierres: {len(filas_a_borrar)} filas eliminadas.")
    return len(filas_a_borrar)


# ── Cierre final vs. parcial ──────────────────────────────────────

COL_ESTADO_CIERRE = "Estado_Cierre"
CIERRE_FINAL = "Final"
CIERRE_PARCIAL = "Parcial"


def anio_en_curso() -> int:
    """Año calendario en curso: el único cuyo cierre es parcial. Todos los anteriores son finales."""
    return date.today().year


def estado_cierre(fecha: object, anio_curso: Optional[int] = None) -> str:
    """'Parcial' si el corte es del año en curso; 'Final' en cualquier año anterior;
    '' si no hay fecha."""
    curso = anio_en_curso() if anio_curso is None else anio_curso
    f = pd.to_datetime(fecha, errors="coerce")
    if pd.isna(f):
        return ""
    return CIERRE_PARCIAL if f.year >= curso else CIERRE_FINAL


def reemplazar_cierres_parciales(
    ws,
    registros: List[Dict],
    anio_curso: Optional[int] = None,
    hoy: Optional[pd.Timestamp] = None,
) -> Tuple[List[Dict], int, int]:
    """Deja UN cierre parcial por Id en el año en curso: el último corte con fecha real.

    1. Entre los registros nuevos del año en curso se elige, por Id, el corte más reciente
       con fecha <= hoy. Un corte con fecha futura (p. ej. un 'diciembre' que repite el
       último valor de un indicador semestral) solo se usa si no hay ninguno real.
    2. Contra la hoja: si el corte elegido es igual o posterior al último corte real
       existente, se borran las filas existentes de ese Id+año (incluidas las de fecha
       futura) y se escribe el nuevo; si es anterior se descarta y se mantiene lo existente.
    Si no llega registro nuevo para el Id+año, la fila existente no se toca.
    Los registros de años cerrados pasan sin cambios.

    Retorna (registros_a_escribir, filas_reemplazadas, nuevos_descartados).
    """
    curso = anio_en_curso() if anio_curso is None else anio_curso
    hoy = pd.Timestamp(date.today()) if hoy is None else pd.Timestamp(hoy)
    cm = _build_col_map(ws)
    idx_fecha = cm.get("Fecha", 6) - 1

    existentes: Dict = defaultdict(list)  # (id, año) -> [(row_idx, fecha)]
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        f = pd.to_datetime(row[idx_fecha].value, errors="coerce")
        if pd.isna(f) or f.year < curso:
            continue
        existentes[(_id_str(row[0].value), f.year)].append((row[0].row, f))

    def _clave(reg: Dict) -> Tuple[bool, pd.Timestamp]:
        f = pd.to_datetime(reg.get("fecha"), errors="coerce")
        return (f <= hoy, f)  # primero los reales (<= hoy), luego el más reciente

    # 1. Un único registro nuevo por Id+año en curso.
    elegidos: Dict = {}
    cerrados: List[Dict] = []
    for reg in registros:
        f = pd.to_datetime(reg.get("fecha"), errors="coerce")
        if pd.isna(f) or f.year < curso:
            cerrados.append(reg)
            continue
        k = (_id_str(reg.get("Id")), f.year)
        if k not in elegidos or _clave(reg) >= _clave(elegidos[k]):
            elegidos[k] = reg
    descartados = sum(
        1
        for reg in registros
        if not pd.isna(pd.to_datetime(reg.get("fecha"), errors="coerce"))
        and pd.to_datetime(reg.get("fecha")).year >= curso
    ) - len(elegidos)

    # 2. Contra lo existente en la hoja.
    a_conservar: List[Dict] = list(cerrados)
    filas_borrar: Set[int] = set()
    for k, reg in elegidos.items():
        f = pd.to_datetime(reg.get("fecha"))
        previas = existentes.get(k, [])
        reales = [fp for _, fp in previas if fp <= hoy]
        if not previas:
            escribir = True
        elif f <= hoy:
            escribir = not reales or f >= max(reales)
        else:  # solo hay un corte futuro (placeholder): nunca pisa un corte real existente
            escribir = not reales
        if escribir:
            filas_borrar.update(r for r, _ in previas)
            a_conservar.append(reg)
        else:
            descartados += 1

    for r_idx in sorted(filas_borrar, reverse=True):
        ws.delete_rows(r_idx)
    logger.info(
        f"  reemplazar_cierres_parciales: {len(filas_borrar)} filas reemplazadas, "
        f"{descartados} registros nuevos descartados."
    )
    return a_conservar, len(filas_borrar), descartados


def marcar_estado_cierre(ws, anio_curso: Optional[int] = None) -> int:
    """Escribe 'Estado_Cierre' (Final/Parcial) en cada fila de Consolidado Cierres.
    Crea la columna al final del encabezado si no existe. Retorna filas marcadas."""
    header = [c for c in next(ws.iter_rows(min_row=1, max_row=1)) if c.value is not None]
    col = next((c.column for c in header if str(c.value).strip() == COL_ESTADO_CIERRE), None)
    if col is None:
        col = max((c.column for c in header), default=0) + 1
        ws.cell(1, col).value = COL_ESTADO_CIERRE
    idx_fecha = _build_col_map(ws).get("Fecha", 6)
    n = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        ws.cell(r, col).value = estado_cierre(ws.cell(r, idx_fecha).value, anio_curso)
        n += 1
    return n


HOJA_MENSUAL = "Consolidado Mensual"


def asegurar_hoja_mensual(wb, ws_ref, nombre: str = HOJA_MENSUAL):
    """Devuelve la hoja de cortes mensuales; si no existe la crea con el mismo encabezado que
    Consolidado Cierres (sin la columna Estado_Cierre, que es solo de Cierres)."""
    if nombre in wb.sheetnames:
        return wb[nombre]
    ws = wb.create_sheet(nombre, index=wb.sheetnames.index(ws_ref.title) + 1)
    col = 0
    for celda in next(ws_ref.iter_rows(min_row=1, max_row=1)):
        if celda.value is None or str(celda.value).strip() == COL_ESTADO_CIERRE:
            continue
        col += 1
        destino = ws.cell(1, col)
        destino.value = celda.value
        if celda.has_style:
            destino.font, destino.fill = celda.font.copy(), celda.fill.copy()
            destino.alignment, destino.border = celda.alignment.copy(), celda.border.copy()
    return ws


def vaciar_hoja(ws) -> int:
    """Borra todas las filas de datos (conserva el encabezado). Retorna cuántas borró."""
    n = max(ws.max_row - 1, 0)
    if n:
        ws.delete_rows(2, n)
    return n


COLOR_DUPLICADO_CIERRE = "F8CBAD"  # naranja claro: Id con más de un cierre en el mismo año


def resaltar_duplicados_id_anio(ws) -> int:
    """Marca con color (sin borrar) las filas de Consolidado Cierres que comparten Id y año.

    Un Id debería tener un único cierre por año; si hay más (p. ej. el mismo Id con dos nombres de
    indicador) se conservan ambos y se resaltan para revisión. Quita primero el resalte anterior,
    así que es idempotente. Retorna cuántas filas quedaron marcadas.
    """
    relleno = PatternFill("solid", fgColor=COLOR_DUPLICADO_CIERRE)
    sin_relleno = PatternFill(fill_type=None)
    idx_fecha = _build_col_map(ws).get("Fecha", 6) - 1
    ultima_col = max((c.column for c in next(ws.iter_rows(min_row=1, max_row=1)) if c.value is not None), default=1)

    filas: Dict = defaultdict(list)  # (id, año) -> [row_idx]
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        for c in row[:ultima_col]:  # limpia el resalte previo (solo el de este color)
            if c.fill is not None and c.fill.fgColor is not None and str(c.fill.fgColor.rgb or "").endswith(COLOR_DUPLICADO_CIERRE):
                c.fill = sin_relleno
        f = pd.to_datetime(row[idx_fecha].value, errors="coerce")
        if not pd.isna(f):
            filas[(_id_str(row[0].value), f.year)].append(row[0].row)

    n = 0
    for rows in filas.values():
        if len(rows) > 1:
            for r in rows:
                for col in range(1, ultima_col + 1):
                    ws.cell(r, col).fill = relleno
                n += 1
    return n


def _dedup_cierres_por_año(ws) -> int:
    """Garantiza UN registro por Id+Año en Consolidado Cierres (el más reciente)."""
    cm = _build_col_map(ws)
    idx_fecha = cm.get("Fecha", 6) - 1
    idx_ejec  = cm.get("Ejecucion", 11) - 1

    filas = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        fecha_raw = row[idx_fecha].value
        try:    fecha = pd.to_datetime(fecha_raw)
        except: fecha = None
        ejec_val = row[idx_ejec].value if len(row) > idx_ejec else None
        filas.append({
            "row_idx": row[0].row,
            "Id":      _id_str(row[0].value),
            "fecha":   fecha,
            "año":     fecha.year if fecha else None,
            "ejec":    ejec_val,
        })

    if not filas:
        return 0

    grupos: Dict = defaultdict(list)
    for f in filas:
        if f["año"] is None:
            continue
        grupos[(f["Id"], f["año"])].append(f)

    filas_a_borrar = []
    for (_id, _año), grupo in grupos.items():
        if len(grupo) <= 1:
            continue
        mejor = max(grupo, key=lambda f: (
            f["fecha"] or pd.Timestamp.min,
            _ejec_score(f["ejec"]),
        ))
        filas_a_borrar.extend(
            f["row_idx"] for f in grupo if f["row_idx"] != mejor["row_idx"]
        )

    for r_idx in sorted(filas_a_borrar, reverse=True):
        ws.delete_rows(r_idx)

    logger.info(
        f"  _dedup_cierres_por_año: {len(filas_a_borrar)} duplicados eliminados."
    )
    return len(filas_a_borrar)


# ── Reparar Meta corrompida por el antiguo capping AGENT5 ─────────

def reparar_meta_capeada_agent5(ws, df_fuente_api, nombre: str = "") -> int:
    """
    Corrige Meta==1 dejado por versiones previas de AGENT5Corrections
    (validate_meta capeaba cualquier Meta > 1.0 a 1.0, asumiendo razón
    0-1; el Meta real de este pipeline es porcentual 0-100 o conteo
    bruto — ver agent5_corrections.py). Esas filas ya existentes no se
    regeneran solas porque su LLAVE ya está en el consolidado.

    Restaura Meta desde el campo crudo 'meta' de la fuente API para esa
    (Id, Fecha). Meta==1 nunca es un valor legítimo en este pipeline,
    así que no hay riesgo de tocar datos correctos.
    """
    lookup: Dict = {}
    for _, r in df_fuente_api.iterrows():
        id_s = _id_str(r.get("Id") or r.get("ID", ""))
        try:
            fecha_key = pd.to_datetime(r["fecha"]).normalize()
        except Exception:
            continue
        meta_raw = pd.to_numeric(r.get("meta"), errors="coerce")
        if pd.notna(meta_raw):
            lookup[(id_s, fecha_key)] = float(meta_raw)

    cm = _build_col_map(ws)
    idx_id    = cm.get("Id")
    idx_fecha = cm.get("Fecha")
    idx_meta  = cm.get("Meta")
    if not all([idx_id, idx_fecha, idx_meta]):
        return 0

    n_fix = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        meta_cell = row[idx_meta - 1]
        try:
            if float(meta_cell.value) != 1.0:
                continue
        except (TypeError, ValueError):
            continue

        id_s = _id_str(row[idx_id - 1].value)
        try:
            fecha_key = pd.to_datetime(row[idx_fecha - 1].value).normalize()
        except Exception:
            continue
        meta_correcta = lookup.get((id_s, fecha_key))
        if meta_correcta is not None and meta_correcta != 1.0:
            meta_cell.value = meta_correcta
            meta_cell.number_format = "General"
            n_fix += 1

    if n_fix:
        logger.info(f"  [{nombre}] Meta corregida (capping AGENT5 obsoleto): {n_fix}")
    else:
        logger.info(f"  [{nombre}] Sin Meta=1 residual de AGENT5.")
    return n_fix


def reparar_ejecucion_capeada_agent5(ws, df_fuente_api, nombre: str = "") -> int:
    """
    Corrige Ejecucion==1.3 dejado por versiones previas de AGENT5Corrections
    (apply_ejecucion_capping capeaba cualquier Ejecucion > 1.3 a 1.3,
    asumiendo razón 0-1.3; el Ejecucion real de este pipeline usa la
    escala propia de cada indicador — porcentual, conteo bruto o incluso
    montos financieros — ver agent5_corrections.py). Esas filas ya
    existentes no se regeneran solas porque su LLAVE ya está en el
    consolidado.

    Restaura Ejecucion desde el campo crudo 'resultado' de la fuente API
    para esa (Id, Fecha). Si el indicador es Promedio/Acumulado, el valor
    correcto (agregado) se recalcula después en reparar_semestral_agregados/
    reparar_multiserie, así que un valor de un solo mes aquí es solo
    temporal y no queda como definitivo.
    """
    lookup: Dict = {}
    for _, r in df_fuente_api.iterrows():
        id_s = _id_str(r.get("Id") or r.get("ID", ""))
        try:
            fecha_key = pd.to_datetime(r["fecha"]).normalize()
        except Exception:
            continue
        ejec_raw = pd.to_numeric(r.get("resultado"), errors="coerce")
        if pd.notna(ejec_raw):
            lookup[(id_s, fecha_key)] = float(ejec_raw)

    cm = _build_col_map(ws)
    idx_id    = cm.get("Id")
    idx_fecha = cm.get("Fecha")
    idx_ejec  = cm.get("Ejecucion")
    if not all([idx_id, idx_fecha, idx_ejec]):
        return 0

    n_fix = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        ejec_cell = row[idx_ejec - 1]
        try:
            if float(ejec_cell.value) != 1.3:
                continue
        except (TypeError, ValueError):
            continue

        id_s = _id_str(row[idx_id - 1].value)
        try:
            fecha_key = pd.to_datetime(row[idx_fecha - 1].value).normalize()
        except Exception:
            continue
        ejec_correcta = lookup.get((id_s, fecha_key))
        if ejec_correcta is not None and ejec_correcta != 1.3:
            ejec_cell.value = ejec_correcta
            ejec_cell.number_format = "General"
            n_fix += 1

    if n_fix:
        logger.info(f"  [{nombre}] Ejecucion corregida (capping AGENT5 obsoleto): {n_fix}")
    else:
        logger.info(f"  [{nombre}] Sin Ejecucion=1.3 residual de AGENT5.")
    return n_fix


# ── Reparar signos desincronizados del catálogo ────────────────────

def reparar_signos_desde_catalogo(
    ws, formato_valores_map: Dict, nombre: str = "",
) -> int:
    """
    Corrige Meta_Signo/Ejecucion_Signo cuando difieren del Formato_Valores
    del catálogo (hallado en 274 y 200+ ids más — el signo históricamente
    solo se heredaba de la última fila real, sin cruzarse jamás contra el
    catálogo, así que quedaba desincronizado cuando el formato real del
    indicador cambiaba o nunca se sincronizó — feedback 2026-07-26).

    No toca filas 'No Aplica' (su signo es un estado de período, no un
    formato de valor) ni ids sin Formato_Valores definido en catálogo.
    Sub-indicadores con Id decimal (274.1) heredan el formato de su padre.
    """
    if not formato_valores_map:
        return 0
    cm = _build_col_map(ws)
    idx_id = cm.get("Id")
    idx_ms = cm.get("MetaS")
    idx_es = cm.get("EjecS")
    if not all([idx_id, idx_ms, idx_es]):
        return 0

    n_fix = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        id_s = _id_str(row[idx_id - 1].value)
        fv = formato_valores_map.get(id_s)
        if not fv:
            padre = id_s.split(".")[0]
            if padre != id_s:
                fv = formato_valores_map.get(padre)
        if not fv:
            continue

        ms_cell = row[idx_ms - 1]
        es_cell = row[idx_es - 1]
        if str(ms_cell.value or "").strip().lower() == "no aplica" \
                or str(es_cell.value or "").strip().lower() == "no aplica":
            continue

        if str(ms_cell.value or "").strip() != fv:
            ms_cell.value = fv
            n_fix += 1
        if str(es_cell.value or "").strip() != fv:
            es_cell.value = fv
            n_fix += 1

    if n_fix:
        logger.info(f"  [{nombre}] Signos corregidos desde catálogo (Formato_Valores): {n_fix}")
    else:
        logger.info(f"  [{nombre}] Signos ya sincronizados con el catálogo.")
    return n_fix


# ── Reparar Meta / Ejecucion vacías ───────────────────────────────

def reparar_meta_vacia(ws, api_kawak_lookup: Dict, nombre: str = "") -> int:
    """Rellena Meta (y Ejecucion) vacías usando el lookup API_Kawak."""
    if not api_kawak_lookup:
        return 0
    cm = _build_col_map(ws)
    idx_id    = cm.get("Id")
    idx_fecha = cm.get("Fecha")
    idx_meta  = cm.get("Meta")
    idx_ejec  = cm.get("Ejecucion")
    idx_tiporeg = cm.get("Tipo_Registro") or cm.get("TipoRegistro")
    idx_ejecs   = cm.get("Ejecucion_Signo") or cm.get("EjecS")
    if not all([idx_id, idx_fecha, idx_meta, idx_ejec]):
        return 0

    n_meta = n_ejec = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        meta_cell = row[idx_meta - 1]
        ejec_cell = row[idx_ejec - 1]
        meta_val  = meta_cell.value
        ejec_val  = ejec_cell.value

        if meta_val is not None and not (isinstance(meta_val, float) and np.isnan(meta_val)):
            continue

        try:
            fecha_key = pd.to_datetime(row[idx_fecha - 1].value).normalize()
        except Exception:
            continue

        id_s = _id_str(row[idx_id - 1].value)
        vals = api_kawak_lookup.get((id_s, fecha_key))
        if vals is None:
            continue

        meta_lookup, ejec_lookup = vals
        if meta_lookup is not None:
            meta_cell.value = meta_lookup
            meta_cell.number_format = "General"
            n_meta += 1
        if (ejec_val is None or (isinstance(ejec_val, float) and np.isnan(ejec_val))) \
                and ejec_lookup is not None:
            ejec_cell.value = ejec_lookup
            n_ejec += 1

    if n_meta or n_ejec:
        logger.info(f"  [{nombre}] Meta reparada: {n_meta} | Ejecucion reparada: {n_ejec}")
    else:
        logger.info(f"  [{nombre}] Sin Meta vacía reparable.")
    return n_meta


def reparar_multiserie(
    ws,
    api_kawak_lookup: Dict,
    tipo_calculo_map: Dict,
    nombre: str = "",
    extraccion_map: Optional[Dict] = None,
    tipo_indicador_map: Optional[Dict] = None,
) -> int:
    """
    Para indicadores multiserie (TipoCalculo definido), sobreescribe Meta
    y Ejecucion con valores del lookup, incluso si ya tienen valor.

    Excluye indicadores 'Desglose Variables' canónicos (se extraen por
    símbolo, no por el campo crudo resultado/meta del lookup — ver
    reparar_desglose_variables) y los de Meta fija por regla de negocio
    (ver _IDS_META_FIJA / reparar_metas_fijas), para no sobreescribirlos
    con el valor crudo equivocado.
    """
    if not api_kawak_lookup or not tipo_calculo_map:
        return 0
    cm = _build_col_map(ws)
    idx_id    = cm.get("Id")
    idx_fecha = cm.get("Fecha")
    idx_meta  = cm.get("Meta")
    idx_ejec  = cm.get("Ejecucion")
    idx_tiporeg = cm.get("Tipo_Registro") or cm.get("TipoRegistro")
    idx_ejecs   = cm.get("Ejecucion_Signo") or cm.get("EjecS")
    if not all([idx_id, idx_fecha, idx_meta, idx_ejec]):
        return 0

    n_meta = n_ejec = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue

        # Mantener intactas las filas sin reporte explícito
        tipo_reg = row[idx_tiporeg - 1].value if idx_tiporeg else None
        ejec_sgn = row[idx_ejecs - 1].value if idx_ejecs else None
        if str(tipo_reg or "").strip().lower() == "no aplica" \
                or str(ejec_sgn or "").strip().lower() == "no aplica":
            continue

        id_s = _id_str(row[idx_id - 1].value)
        if id_s not in tipo_calculo_map:
            continue
        if _usa_variables_canonico(id_s, (extraccion_map or {}).get(id_s), tipo_indicador_map):
            continue
        if id_s in _IDS_SUMA_VARIABLES_SERIES:
            continue
        try:
            fecha_key = pd.to_datetime(row[idx_fecha - 1].value).normalize()
        except Exception:
            continue
        vals = api_kawak_lookup.get((id_s, fecha_key))
        if vals is None:
            continue

        meta_cell = row[idx_meta - 1]
        ejec_cell = row[idx_ejec - 1]
        meta_lookup, ejec_lookup = vals

        # Meta fija por regla de negocio (ver _IDS_META_FIJA): nunca se
        # sobreescribe con el valor crudo del lookup, sólo con reparar_metas_fijas.
        if id_s not in _IDS_META_FIJA:
            if meta_lookup is not None and meta_cell.value != meta_lookup:
                meta_cell.value = meta_lookup
                meta_cell.number_format = "General"
                n_meta += 1
        if ejec_lookup is not None and ejec_cell.value != ejec_lookup:
            ejec_cell.value = ejec_lookup
            n_ejec += 1

    if n_meta or n_ejec:
        logger.info(f"  [{nombre}] Multiserie corregida: Meta={n_meta} Ejec={n_ejec}")
    else:
        logger.info(f"  [{nombre}] Multiserie: sin correcciones.")
    return n_meta


def reparar_desglose_variables(
    ws,
    df_fuente_api: pd.DataFrame,
    extraccion_map: Dict,
    variables_campo_map: Dict,
    tipo_indicador_map: Dict,
    nombre: str = "",
    tipo_calculo_map: Optional[Dict] = None,
) -> int:
    """
    Para indicadores 'Desglose Variables' canónicos, recalcula Meta y
    Ejecucion desde los símbolos de la hoja Variables — corrige el efecto
    de reparar_multiserie (versiones previas del pipeline), que sobreescribía
    con el campo crudo resultado/meta en vez del símbolo real.

    Si se pasa tipo_calculo_map, excluye los IDs Promedio/Acumulado: esos
    se agregan por período en reparar_semestral_agregados y no deben
    pisarse aquí con el valor de un único mes.
    """
    ids_canonico = {
        ids for ids, ext in (extraccion_map or {}).items()
        if _usa_variables_canonico(ids, ext, tipo_indicador_map)
    }
    if tipo_calculo_map:
        ids_agg = {
            ids for ids, tc in tipo_calculo_map.items()
            if str(tc).strip().lower() in ("promedio", "acumulado")
        }
        ids_canonico -= ids_agg
    if not ids_canonico:
        return 0

    lookup: Dict = {}
    for _, r in df_fuente_api.iterrows():
        id_s = _id_str(r.get("Id") or r.get("ID", ""))
        if id_s not in ids_canonico:
            continue
        try:
            fecha_key = pd.to_datetime(r["fecha"]).normalize()
        except Exception:
            continue
        ejec_v = _ejec_corrected_from_row(
            r, extraccion_map, None, variables_campo_map, tipo_indicador_map
        )
        meta_v = _meta_corrected_from_row(
            r, extraccion_map, None, variables_campo_map, tipo_indicador_map
        )
        lookup[(id_s, fecha_key)] = (ejec_v, meta_v)

    cm = _build_col_map(ws)
    idx_id    = cm.get("Id")
    idx_fecha = cm.get("Fecha")
    idx_meta  = cm.get("Meta")
    idx_ejec  = cm.get("Ejecucion")
    idx_tiporeg = cm.get("Tipo_Registro") or cm.get("TipoRegistro")
    idx_ejecs   = cm.get("Ejecucion_Signo") or cm.get("EjecS")
    if not all([idx_id, idx_fecha, idx_meta, idx_ejec]):
        return 0

    n_meta = n_ejec = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        tipo_reg = row[idx_tiporeg - 1].value if idx_tiporeg else None
        ejec_sgn = row[idx_ejecs - 1].value if idx_ejecs else None
        if str(tipo_reg or "").strip().lower() == "no aplica" \
                or str(ejec_sgn or "").strip().lower() == "no aplica":
            continue
        id_s = _id_str(row[idx_id - 1].value)
        if id_s not in ids_canonico:
            continue
        try:
            fecha_key = pd.to_datetime(row[idx_fecha - 1].value).normalize()
        except Exception:
            continue
        vals = lookup.get((id_s, fecha_key))
        if vals is None:
            continue
        ejec_v, meta_v = vals

        meta_cell = row[idx_meta - 1]
        ejec_cell = row[idx_ejec - 1]
        if meta_v is not None and meta_cell.value != meta_v:
            meta_cell.value = meta_v
            meta_cell.number_format = "General"
            n_meta += 1
        if ejec_v is not None and ejec_cell.value != ejec_v:
            ejec_cell.value = ejec_v
            n_ejec += 1

    if n_meta or n_ejec:
        logger.info(f"  [{nombre}] Desglose Variables corregido: Meta={n_meta} Ejec={n_ejec}")
    else:
        logger.info(f"  [{nombre}] Desglose Variables: sin correcciones.")
    return n_meta + n_ejec


def reparar_metas_fijas(ws, nombre: str = "") -> int:
    """
    Fuerza la Meta de indicadores con regla de negocio de meta fija
    (ver _IDS_META_FIJA) en TODAS las filas existentes, incluso si ya
    tenían un valor distinto — última pasada, siempre gana.
    """
    if not _IDS_META_FIJA:
        return 0
    cm = _build_col_map(ws)
    idx_id   = cm.get("Id")
    idx_meta = cm.get("Meta")
    if not all([idx_id, idx_meta]):
        return 0

    n = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        id_s = _id_str(row[idx_id - 1].value)
        meta_fija = _IDS_META_FIJA.get(id_s)
        if meta_fija is None:
            continue
        meta_cell = row[idx_meta - 1]
        if meta_cell.value != meta_fija:
            meta_cell.value = meta_fija
            meta_cell.number_format = "General"
            n += 1

    if n:
        logger.info(f"  [{nombre}] Metas fijas forzadas: {n}")
    return n


def reparar_semestral_agregados(
    ws,
    df_fuente_api: pd.DataFrame,
    extraccion_map: Dict,
    tipo_calculo_map: Dict,
    nombre: str = "",
    variables_campo_map: Optional[Dict] = None,
    tipo_indicador_map: Optional[Dict] = None,
) -> int:
    """
    Para indicadores Promedio/Acumulado, recalcula Meta y Ejecucion
    desde los datos mensuales (corrige el efecto de reparar_multiserie
    que escribe el valor mensual en lugar del agregado semestral/anual).
    """
    ids_avg = {ids for ids, tc in tipo_calculo_map.items() if tc.lower() == "promedio"}
    ids_sum = {ids for ids, tc in tipo_calculo_map.items() if tc.lower() == "acumulado"}
    ids_agg = ids_avg | ids_sum
    if not ids_agg:
        logger.info(f"  [{nombre}] Sin indicadores Promedio/Acumulado.")
        return 0

    # Lookup mensual: {(id_s, year, month): (ejec, meta)}
    monthly: Dict = {}
    for _, r in df_fuente_api.iterrows():
        id_s = _id_str(r.get("Id") or r.get("ID", ""))
        if id_s not in ids_agg:
            continue
        try:
            fecha = pd.to_datetime(r["fecha"])
        except Exception:
            continue
        ejec_v = _ejec_corrected_from_row(
            r, extraccion_map, None, variables_campo_map, tipo_indicador_map
        )
        meta_v = _meta_corrected_from_row(
            r, extraccion_map, None, variables_campo_map, tipo_indicador_map
        )
        monthly[(id_s, fecha.year, fecha.month)] = (ejec_v, meta_v)

    cm = _build_col_map(ws)
    idx_id    = cm.get("Id")
    idx_fecha = cm.get("Fecha")
    idx_meta  = cm.get("Meta")
    idx_ejec  = cm.get("Ejecucion")
    idx_tiporeg = cm.get("Tipo_Registro") or cm.get("TipoRegistro")
    idx_ejecs   = cm.get("Ejecucion_Signo") or cm.get("EjecS")
    if not all([idx_id, idx_fecha, idx_meta, idx_ejec]):
        return 0

    is_cierre = "Cierre" in ws.title or "Cierres" in ws.title
    is_mensual = "Mensual" in ws.title  # corte acumulado: enero hasta el mes de la fila
    n_fix = 0

    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue

        # Si el periodo está marcado como No Aplica, no recalcular agregados
        tipo_reg = row[idx_tiporeg - 1].value if idx_tiporeg else None
        ejec_sgn = row[idx_ejecs - 1].value if idx_ejecs else None
        if str(tipo_reg or "").strip().lower() == "no aplica" \
                or str(ejec_sgn or "").strip().lower() == "no aplica":
            continue

        id_s = _id_str(row[idx_id - 1].value)
        if id_s not in ids_agg:
            continue
        try:
            fecha = pd.to_datetime(row[idx_fecha - 1].value)
        except Exception:
            continue

        patron = "AVG" if id_s in ids_avg else "SUM"
        if is_mensual:
            months = list(range(1, fecha.month + 1))
        elif is_cierre:
            months = list(range(1, 13))
        else:
            months = list(range(1, 7)) if fecha.month <= 6 else list(range(7, 13))

        pares = [monthly.get((id_s, fecha.year, m), (None, None)) for m in months]
        if patron == "AVG":
            # Meta=0 y Ejecución=0 en un mes del período: sin medición, no entra al promedio
            # (solo si hay otros meses con dato; si todo es 0, el promedio sigue siendo 0).
            con_dato = [(e, mt) for e, mt in pares if not (e == 0 and mt == 0)]
            if any(e is not None for e, _ in con_dato):
                pares = con_dato
        ejecs = [e for e, _ in pares if e is not None]
        metas = [mt for _, mt in pares if mt is not None]

        if not ejecs:
            # Sin dato mensual real para este período: si el valor actual es
            # exactamente 1.3, es un residuo del capping AGENT5 obsoleto
            # (ningún dato lo respalda) — se deja en blanco en vez de
            # conservar una cifra fabricada.
            ejec_cell = row[idx_ejec - 1]
            try:
                if float(ejec_cell.value) == 1.3:
                    ejec_cell.value = None
                    n_fix += 1
            except (TypeError, ValueError):
                pass
            continue

        ejec_agg = sum(ejecs) / len(ejecs) if patron == "AVG" else sum(ejecs)
        meta_agg = (
            (sum(metas) / len(metas) if patron == "AVG" else sum(metas))
            if metas else None
        )

        ejec_cell = row[idx_ejec - 1]
        meta_cell = row[idx_meta - 1]
        if ejec_cell.value != ejec_agg:
            ejec_cell.value = ejec_agg
            n_fix += 1
        if meta_agg is not None and meta_cell.value != meta_agg:
            meta_cell.value = meta_agg
            meta_cell.number_format = "General"

    logger.info(f"  [{nombre}] Promedio/Acumulado recalculados: {n_fix} filas")
    return n_fix
