"""Constructores de datos para CMI por Procesos — paridad con Streamlit resumen_por_proceso."""

from __future__ import annotations

import math
import re
import unicodedata
from datetime import date
from typing import Any

import pandas as pd

from app.domain.categorization import categorizar_cumplimiento
from app.domain.cmi_builders import (
    build_distribucion_nivel,
    calcular_kpis,
    df_to_records,
)
from app.domain.constants import COLOR_CATEGORIA
from app.domain.resumen_builders import compute_trends

MESES_OPCIONES = [
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
]

MES_MAP: dict[str, int] = {m.upper(): i + 1 for i, m in enumerate(MESES_OPCIONES)}

TIPOS_PROCESO = ["APOYO", "ESTRATÉGICO", "EVALUACIÓN Y MEJORA", "MISIONAL"]

TIPO_PROCESO_COLORS: dict[str, str] = {
    "APOYO": "#42A5F5",
    "ESTRATÉGICO": "#1A3A5C",
    "EVALUACIÓN Y MEJORA": "#FB8C00",
    "MISIONAL": "#43A047",
}

TIPO_PROCESO_COLORS_LIGHT: dict[str, str] = {
    "APOYO": "#E3F2FD",
    "ESTRATÉGICO": "#D0E4FF",
    "EVALUACIÓN Y MEJORA": "#FFF3E0",
    "MISIONAL": "#E8F5E9",
}

TIPO_PROCESO_ICONS: dict[str, str] = {
    "APOYO": "🔧",
    "ESTRATÉGICO": "🎯",
    "EVALUACIÓN Y MEJORA": "📊",
    "MISIONAL": "⭐",
}

STRATEGIC_PALETTE = ["#FBAF17", "#42F2F2", "#EC0677", "#1FB2DE", "#A6CE38", "#0F385A"]


_ICONO_CATEGORIA: dict[str, str] = {
    "Sobrecumplimiento": "🟢",
    "Cumplimiento": "🟢",
    "Alerta": "🟡",
    "Peligro": "🔴",
    "Sin dato": "—",
}


def cumplimiento_semaforo_color(val: float | None) -> str:
    """Color por régimen general (agregados de proceso, sin Id único) —
    delega a categorizar_cumplimiento, fuente única de verdad. Corregido en
    Oleada 2: antes usaba umbrales propios (100/80), divergentes del
    régimen general canónico (100/105 para Sobrecumplimiento) — ver
    docs/tecnico/05-reglas-de-negocio.md."""
    categoria = categorizar_cumplimiento(val / 100.0 if val is not None else None)
    return COLOR_CATEGORIA[categoria]


def cumplimiento_estado(val: float | None) -> dict[str, str]:
    """Ídem cumplimiento_semaforo_color, con label/icon — vocabulario
    canónico (Peligro/Alerta/Cumplimiento/Sobrecumplimiento/Sin dato) en
    vez del propio (Saludable/Crítico) que tenía antes."""
    categoria = categorizar_cumplimiento(val / 100.0 if val is not None else None)
    return {
        "label": categoria,
        "icon": _ICONO_CATEGORIA[categoria],
        "color": COLOR_CATEGORIA[categoria],
    }


def _norm_text(value: object) -> str:
    txt = str(value or "").strip().upper()
    txt = unicodedata.normalize("NFD", txt)
    return "".join(c for c in txt if unicodedata.category(c) != "Mn")


def _safe_float(val: Any, default: float | None = None) -> float | None:
    if val is None:
        return default
    try:
        f = float(val)
        if math.isnan(f) or math.isinf(f):
            return default
        return f
    except (TypeError, ValueError):
        return default


def mes_to_num(value: object, periodicidad: object = None) -> float | None:
    if pd.isna(value):
        return None
    period_norm = str(periodicidad or "").strip().lower()
    if isinstance(value, int | float):
        try:
            v = int(value)
            if period_norm == "semestral":
                if v in (1, 2):
                    return 6.0
                if v == 7:
                    return 12.0
            return float(v) if 1 <= v <= 12 else None
        except Exception:
            return None
    txt = str(value).strip()
    if not txt:
        return None
    txt_norm = _norm_text(txt)
    if period_norm == "semestral":
        if txt_norm in ("1", "01", "ENERO", "ENE"):
            return 6.0
        if txt_norm in ("2", "02", "JULIO", "JUL", "7", "07"):
            return 12.0
    if txt.isdigit():
        v = int(txt)
        return float(v) if 1 <= v <= 12 else None
    return float(MES_MAP[txt_norm]) if txt_norm in MES_MAP else None


def mes_nombre(mes_num: int) -> str:
    if 1 <= mes_num <= 12:
        return MESES_OPCIONES[mes_num - 1]
    return "Diciembre"


def default_mes(tracking: pd.DataFrame, anio: int) -> int:
    prev = get_prev_month_for_year(tracking, anio)
    return prev if prev is not None else 12


def get_prev_month_for_year(tracking_df: pd.DataFrame, year: int) -> int | None:
    if tracking_df.empty or "Anio" not in tracking_df.columns:
        return None
    subset = tracking_df[pd.to_numeric(tracking_df["Anio"], errors="coerce") == int(year)].copy()
    if subset.empty or "Mes" not in subset.columns:
        return None
    nums = subset["Mes"].apply(lambda m: mes_to_num(m)).dropna()
    if nums.empty:
        return None
    return int(nums.max())


def _ensure_cumplimiento_pct(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "cumplimiento_pct" not in out.columns:
        if "Cumplimiento_norm" in out.columns:
            out["cumplimiento_pct"] = pd.to_numeric(out["Cumplimiento_norm"], errors="coerce") * 100
        elif "Cumplimiento" in out.columns:
            vals = pd.to_numeric(out["Cumplimiento"], errors="coerce")
            if not vals.dropna().empty and vals.max() <= 2:
                vals = vals * 100
            out["cumplimiento_pct"] = vals
    return out


def _id_base(valor: object) -> str:
    s = str(valor).strip()
    return s[:-2] if s.endswith(".0") else s


def aplicar_jerarquia_por_id(
    out: pd.DataFrame, jerarquia: pd.DataFrame, map_df: pd.DataFrame
) -> pd.DataFrame:
    """Asigna Proceso_padre / Subproceso_final por Id según el catálogo maestro.

    El indicador se vincula por Id (un subindicador '386.1' hereda el de '386'); el nombre
    del subproceso solo sirve para traer sus atributos (Unidad, Tipo de proceso) del mapa.
    Las filas cuyo Id no está en el catálogo conservan la asignación previa.
    """
    if out.empty or "Id" not in out.columns or jerarquia is None or jerarquia.empty:
        return out
    jer = jerarquia.drop_duplicates(subset=["Id"]).set_index("Id")
    ids = out["Id"].map(_id_base)
    bases = ids.str.split(".").str[0]

    def _buscar(col: str) -> pd.Series:
        directo = ids.map(jer[col]).replace("", pd.NA)
        return directo.where(directo.notna(), bases.map(jer[col]).replace("", pd.NA))

    def _espacios(v: object) -> object:
        return " ".join(v.replace(" ", " ").split()) if isinstance(v, str) else v

    sub_cat = _buscar("Subproceso_cat").map(_espacios).replace("", pd.NA)
    proc_cat = _buscar("Proceso_cat").map(_espacios).replace("", pd.NA)
    sub_cat = sub_cat.fillna(proc_cat)
    proc_cat = proc_cat.fillna(sub_cat)
    hay = sub_cat.notna()
    if not hay.any():
        return out

    attrs: dict[str, dict[str, str]] = {}
    if not map_df.empty and {"Subproceso", "Proceso"}.issubset(map_df.columns):
        for _, r in map_df.dropna(subset=["Subproceso"]).iterrows():
            attrs.setdefault(
                _norm_text(r["Subproceso"]),
                {
                    "sub": str(r["Subproceso"]).strip(),
                    "proceso": str(r["Proceso"]).strip(),
                    "unidad": str(r.get("Unidad", "") or ""),
                    "tipo": str(r.get("Tipo de proceso", "") or ""),
                },
            )
    claves = sub_cat.map(lambda s: _norm_text(s) if isinstance(s, str) else "")
    info = claves.map(attrs)
    out = out.copy()
    out["TipoCalculo"] = _buscar("TipoCalculo").fillna("")
    en_mapa = hay & info.notna()
    out.loc[hay, "Proceso_padre"] = proc_cat[hay]
    out.loc[hay, "Subproceso_final"] = sub_cat[hay]
    for destino, clave in (
        ("Proceso_padre", "proceso"),
        ("Subproceso_final", "sub"),
        ("Unidad", "unidad"),
        ("Tipo de proceso", "tipo"),
    ):
        if en_mapa.any():
            if destino not in out.columns:
                out[destino] = ""
            out.loc[en_mapa, destino] = info[en_mapa].map(lambda d, k=clave: d[k])
    return out


def prepare_tracking(
    df: pd.DataFrame, map_df: pd.DataFrame, jerarquia: pd.DataFrame | None = None
) -> pd.DataFrame:
    """Enriquece tracking con Proceso_padre, Subproceso_final, Unidad y Tipo de proceso.

    Con `jerarquia` (catálogo maestro) el vínculo indicador→proceso/subproceso se hace por Id.
    """
    if df.empty:
        return df

    out = df.copy()
    if "Proceso" not in out.columns:
        out["Proceso"] = "Sin proceso"

    if not map_df.empty and {"Subproceso", "Proceso"}.issubset(map_df.columns):
        sub_map = (
            map_df[["Subproceso", "Proceso"] + (["Unidad"] if "Unidad" in map_df.columns else [])]
            .dropna(subset=["Subproceso"])
            .drop_duplicates(subset=["Subproceso"])
            .copy()
        )
        sub_map["sub_norm"] = sub_map["Subproceso"].astype(str).map(_norm_text)

        proc_col = "Subproceso" if "Subproceso" in out.columns else "Proceso"
        out["proc_norm"] = out[proc_col].astype(str).map(_norm_text)
        out = out.merge(
            sub_map.rename(columns={"Proceso": "Proceso_padre_map", "Unidad": "Unidad_map"}),
            left_on="proc_norm",
            right_on="sub_norm",
            how="left",
        )
        out["Proceso_padre"] = out["Proceso_padre_map"].fillna(out["Proceso"].astype(str))
        out["Subproceso_final"] = out.get("Subproceso", out["Proceso"]).fillna(
            out["Proceso"].astype(str)
        )
        if "Unidad_map" in out.columns:
            out["Unidad"] = out["Unidad_map"].fillna("")
        out = out.drop(
            columns=[
                c
                for c in ["proc_norm", "sub_norm", "Proceso_padre_map", "Unidad_map"]
                if c in out.columns
            ]
        )
    else:
        out["Proceso_padre"] = out["Proceso"].astype(str)
        out["Subproceso_final"] = out.get("Subproceso", out["Proceso"]).astype(str)
        if "Unidad" not in out.columns:
            out["Unidad"] = ""

    if not map_df.empty and "Tipo de proceso" in map_df.columns and "Subproceso" in map_df.columns:
        tipo_map = map_df[["Subproceso", "Tipo de proceso"]].drop_duplicates(subset=["Subproceso"])
        sub_col = "Subproceso_final" if "Subproceso_final" in out.columns else "Subproceso"
        if sub_col in out.columns:
            out = out.merge(
                tipo_map.rename(columns={"Subproceso": "sub_merge_key"}),
                left_on=sub_col,
                right_on="sub_merge_key",
                how="left",
            )
            if "Tipo de proceso" not in out.columns and "Tipo de proceso_y" in out.columns:
                out["Tipo de proceso"] = out["Tipo de proceso_y"]
            out = out.drop(
                columns=[c for c in ["sub_merge_key", "Tipo de proceso_y"] if c in out.columns]
            )

    if jerarquia is not None:
        out = aplicar_jerarquia_por_id(out, jerarquia, map_df)
    return _ensure_cumplimiento_pct(out)


def filter_by_anio_mes(df: pd.DataFrame, *, anio: int, mes: int) -> pd.DataFrame:
    if df.empty:
        return df
    out = df.copy()
    if "Anio" in out.columns:
        out = out[pd.to_numeric(out["Anio"], errors="coerce") == int(anio)]
    if "Mes" in out.columns and mes:
        out = out[out["Mes"].apply(mes_to_num) == float(mes)]
    return out


def apply_ui_filters(
    df: pd.DataFrame,
    *,
    unidad: str | None = None,
    proceso: str | None = None,
    subproceso: str | None = None,
    clasificacion: str | None = None,
    frecuencia: str | None = None,
    ids_linea: set[str] | None = None,
) -> pd.DataFrame:
    """`ids_linea` (conjunto vacío incluido) restringe a los indicadores de la línea
    estratégica elegida del PDI; None = sin filtro por línea."""
    if df.empty:
        return df
    out = df.copy()
    if ids_linea is not None and "Id" in out.columns:
        out = out[out["Id"].astype(str).str.strip().isin(ids_linea)]
    if unidad and unidad != "Todos" and "Unidad" in out.columns:
        out = out[out["Unidad"].astype(str) == unidad]
    if proceso and proceso != "Todos":
        col = "Proceso_padre" if "Proceso_padre" in out.columns else "Proceso"
        out = out[out[col].astype(str) == proceso]
    if subproceso and subproceso != "Todos":
        col = "Subproceso_final" if "Subproceso_final" in out.columns else "Subproceso"
        if col in out.columns:
            out = out[out[col].astype(str) == subproceso]
    if clasificacion and clasificacion != "Todos" and "Clasificacion" in out.columns:
        out = out[out["Clasificacion"].astype(str) == clasificacion]
    if frecuencia and frecuencia != "Todos":
        freq_col = next((c for c in ["Frecuencia", "Periodicidad"] if c in out.columns), None)
        if freq_col:
            out = out[out[freq_col].astype(str) == frecuencia]
    return out


def latest_per_indicator(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    out = df.copy()
    if "Id" in out.columns:
        return out.drop_duplicates(subset=["Id"], keep="last")
    if "Indicador" in out.columns:
        return out.drop_duplicates(subset=["Indicador"], keep="last")
    return out


def build_procesos_kpis(df: pd.DataFrame) -> dict[str, Any]:
    base = calcular_kpis(df)
    base["n_procesos"] = int(df["Proceso_padre"].nunique()) if "Proceso_padre" in df.columns else 0
    base["n_subprocesos"] = (
        int(df["Subproceso_final"].nunique()) if "Subproceso_final" in df.columns else 0
    )
    base["n_unidades"] = (
        int(df["Unidad"].replace("", pd.NA).dropna().nunique()) if "Unidad" in df.columns else 0
    )
    return base


def build_banner(
    df: pd.DataFrame,
    *,
    anio: int,
    mes: int,
    base_year: int,
    base_month: int | None,
    cumpl_global: float | None,
    cumpl_base: float | None,
) -> dict[str, Any]:
    en_riesgo = 0
    if not df.empty and "Nivel de cumplimiento" in df.columns:
        en_riesgo = int(df["Nivel de cumplimiento"].isin(["Peligro", "Alerta"]).sum())
    variacion = None
    if cumpl_global is not None and cumpl_base is not None:
        variacion = round(cumpl_global - cumpl_base, 1)
    pct_saludable = None
    if not df.empty and "cumplimiento_pct" in df.columns:
        valores = pd.to_numeric(df["cumplimiento_pct"], errors="coerce").dropna()
        if len(valores):
            pct_saludable = round(float((valores >= 100).sum()) / len(valores) * 100, 1)
    return {
        "titulo": "CMI por Procesos",
        "anio": anio,
        "mes": mes_nombre(mes),
        "cumplimiento_global": cumpl_global,
        "cumplimiento_base_anio": cumpl_base,
        "base_anio": base_year,
        "base_mes": mes_nombre(base_month) if base_month else None,
        "variacion_pp": variacion,
        "total_indicadores": len(df),
        "en_riesgo": en_riesgo,
        "pct_saludable": pct_saludable,
    }


AMBIENTAL_PATTERN = re.compile(
    r"ambiental|medio ambiente|sostenibilidad|sustentabilidad", re.IGNORECASE
)


def build_brecha_ambiental(df: pd.DataFrame) -> dict[str, Any]:
    """Paridad con Streamlit: insight de 'Brecha ambiental crítica' del corte actual.

    Filtra indicadores cuyo proceso coincide con temática ambiental/sostenibilidad
    y reporta su cumplimiento promedio, o un mensaje por defecto si no hay datos.
    """
    default_text = "No se identificó una brecha ambiental crítica clara en los indicadores activos."
    if df.empty:
        return {"n_indicadores": 0, "cumplimiento_promedio": None, "texto": default_text}

    col = "Proceso_padre" if "Proceso_padre" in df.columns else "Proceso"
    if col not in df.columns:
        return {"n_indicadores": 0, "cumplimiento_promedio": None, "texto": default_text}

    mask = df[col].astype(str).str.contains(AMBIENTAL_PATTERN, na=False)
    env_df = df[mask]
    if env_df.empty or "cumplimiento_pct" not in env_df.columns:
        return {"n_indicadores": 0, "cumplimiento_promedio": None, "texto": default_text}

    valores = pd.to_numeric(env_df["cumplimiento_pct"], errors="coerce").dropna()
    if valores.empty:
        return {"n_indicadores": 0, "cumplimiento_promedio": None, "texto": default_text}

    count = int(len(valores))
    promedio = round(float(valores.mean()), 1)
    texto = f"Los {count} indicadores de Gestión Ambiental reportan {promedio}% de cumplimiento en el periodo actual."
    return {"n_indicadores": count, "cumplimiento_promedio": promedio, "texto": texto}


def _normalize_tipo(tipo: str) -> str:
    t = _norm_text(tipo)
    if "EVALUACION" in t or "EVALUACIÓN" in t:
        return "EVALUACIÓN Y MEJORA"
    for known in TIPOS_PROCESO:
        if _norm_text(known) == t:
            return known
    return str(tipo or "Sin tipo").strip()


def build_tipo_proceso_cards(
    df: pd.DataFrame, df_prev: pd.DataFrame | None = None
) -> list[dict[str, Any]]:
    if df.empty or "Tipo de proceso" not in df.columns:
        return []
    cards = []
    for tipo_raw, group in df.groupby("Tipo de proceso", dropna=False):
        tipo = _normalize_tipo(str(tipo_raw))
        cumpl = (
            _safe_float(group["cumplimiento_pct"].mean())
            if "cumplimiento_pct" in group.columns
            else None
        )
        prev_cumpl = None
        variacion = None
        if df_prev is not None and not df_prev.empty and "Tipo de proceso" in df_prev.columns:
            prev_group = df_prev[df_prev["Tipo de proceso"].astype(str) == str(tipo_raw)]
            if not prev_group.empty:
                prev_cumpl = _safe_float(prev_group["cumplimiento_pct"].mean())
                if cumpl is not None and prev_cumpl is not None:
                    variacion = round(cumpl - prev_cumpl, 1)
        riesgo = 0
        if "Nivel de cumplimiento" in group.columns:
            riesgo = int(group["Nivel de cumplimiento"].isin(["Peligro", "Alerta"]).sum())
        cards.append(
            {
                "tipo": tipo,
                "color": TIPO_PROCESO_COLORS.get(tipo, "#BDBDBD"),
                "color_light": TIPO_PROCESO_COLORS_LIGHT.get(tipo, "#F5F5F5"),
                "icon": TIPO_PROCESO_ICONS.get(tipo, "📋"),
                "cumplimiento": round(cumpl, 1) if cumpl is not None else None,
                "cumplimiento_anterior": round(prev_cumpl, 1) if prev_cumpl is not None else None,
                "variacion_pp": variacion,
                "n_indicadores": len(group),
                "n_riesgo": riesgo,
            }
        )
    cards.sort(key=lambda x: x["cumplimiento"] or 0, reverse=True)
    return cards


def build_proceso_bars(
    df: pd.DataFrame, df_prev: pd.DataFrame | None = None
) -> list[dict[str, Any]]:
    col = "Proceso_padre" if "Proceso_padre" in df.columns else "Proceso"
    if df.empty or col not in df.columns:
        return []
    bars = []
    for proceso, group in df.groupby(col, dropna=True):
        if not proceso or str(proceso).strip() in ("", "nan"):
            continue
        actual = (
            _safe_float(group["cumplimiento_pct"].mean())
            if "cumplimiento_pct" in group.columns
            else None
        )
        anterior = None
        if df_prev is not None and not df_prev.empty and col in df_prev.columns:
            prev_g = df_prev[df_prev[col].astype(str) == str(proceso)]
            if not prev_g.empty:
                anterior = _safe_float(prev_g["cumplimiento_pct"].mean())
        nivel = None
        if "Nivel de cumplimiento" in group.columns:
            modas = group["Nivel de cumplimiento"].dropna()
            if not modas.empty:
                nivel = modas.mode().iloc[0]
        bars.append(
            {
                "proceso": str(proceso),
                "cumplimiento": round(actual, 1) if actual is not None else None,
                "cumplimiento_anterior": round(anterior, 1) if anterior is not None else None,
                "n_indicadores": len(group),
                "color": cumplimiento_semaforo_color(actual),
                "nivel": nivel,
            }
        )
    bars.sort(key=lambda x: x["cumplimiento"] or 0)
    return bars


def build_procesos_detalle(df: pd.DataFrame) -> list[dict[str, Any]]:
    col = "Proceso_padre" if "Proceso_padre" in df.columns else "Proceso"
    if df.empty or col not in df.columns:
        return []
    items = []
    for proceso, group in df.groupby(col, dropna=True):
        if not proceso or str(proceso).strip() in ("", "nan"):
            continue
        cumpl = _safe_float(group["cumplimiento_pct"].mean())
        conteo: dict[str, int] = {}
        if "Nivel de cumplimiento" in group.columns:
            for cat in ("Peligro", "Alerta", "Cumplimiento", "Sobrecumplimiento"):
                conteo[cat] = int((group["Nivel de cumplimiento"] == cat).sum())
        subprocesos = []
        sub_col = "Subproceso_final" if "Subproceso_final" in group.columns else "Subproceso"
        if sub_col in group.columns:
            for sub, sg in group.groupby(sub_col, dropna=True):
                if not sub or str(sub).strip() in ("", "nan"):
                    continue
                sub_cumpl = _safe_float(sg["cumplimiento_pct"].mean())
                sub_riesgo = (
                    int(sg["Nivel de cumplimiento"].isin(["Peligro", "Alerta"]).sum())
                    if "Nivel de cumplimiento" in sg.columns
                    else 0
                )
                subprocesos.append(
                    {
                        "subproceso": str(sub),
                        "cumplimiento": sub_cumpl,
                        "n_indicadores": len(sg),
                        "en_riesgo": sub_riesgo,
                    }
                )
            subprocesos.sort(key=lambda x: x["cumplimiento"] or 0)
        items.append(
            {
                "proceso": str(proceso),
                "unidad": str(group["Unidad"].iloc[0])
                if "Unidad" in group.columns and len(group)
                else "",
                "tipo_proceso": str(group["Tipo de proceso"].iloc[0])
                if "Tipo de proceso" in group.columns and len(group)
                else "",
                "total_indicadores": len(group),
                "cumplimiento_promedio": cumpl,
                "categorias": conteo,
                "en_riesgo": conteo.get("Peligro", 0) + conteo.get("Alerta", 0),
                "subprocesos": subprocesos,
            }
        )
    items.sort(key=lambda x: x["cumplimiento_promedio"] or 0)
    return items


def build_unidades_detalle(df: pd.DataFrame) -> list[dict[str, Any]]:
    if df.empty or "Unidad" not in df.columns:
        return []
    items = []
    for unidad, group in df.groupby("Unidad", dropna=True):
        u = str(unidad).strip()
        if not u:
            continue
        cumpl = _safe_float(group["cumplimiento_pct"].mean())
        riesgo = (
            int(group["Nivel de cumplimiento"].isin(["Peligro", "Alerta"]).sum())
            if "Nivel de cumplimiento" in group.columns
            else 0
        )
        criticos = (
            int((group["Nivel de cumplimiento"] == "Peligro").sum())
            if "Nivel de cumplimiento" in group.columns
            else 0
        )
        procesos = int(group["Proceso_padre"].nunique()) if "Proceso_padre" in group.columns else 0
        estado = cumplimiento_estado(cumpl)
        items.append(
            {
                "unidad": u,
                "cumplimiento_promedio": cumpl,
                "n_indicadores": len(group),
                "n_procesos": procesos,
                "en_riesgo": riesgo,
                "n_criticos": criticos,
                "color": cumplimiento_semaforo_color(cumpl),
                "estado": estado["label"],
                "estado_icon": estado["icon"],
                "estado_color": estado["color"],
            }
        )
    items.sort(key=lambda x: x["cumplimiento_promedio"] or 0, reverse=True)
    return items


def build_catalog_charts(
    cmi_df: pd.DataFrame, active_ids: set[str]
) -> dict[str, list[dict[str, Any]]]:
    """KPIs de catálogo para CMI Procesos: periodicidad y clasificación.

    Nota (2026-09-18, ver docs/migration/PLAN_MIGRACION_PRIORIZADO.md): la
    columna "Tipo de indicador" del Excel ("Tipo 1"/"Tipo 2") es un flag
    técnico interno del pipeline de extracción (variables vs. series), NO una
    categoría de negocio — nunca se muestra al usuario en el legacy. Aquí se
    usa "Clasificacion" (Estratégico/Operativo), que sí es la categoría real
    y ya se usa como filtro visible en el resto de este módulo.
    """
    periodicidad: list[dict[str, Any]] = []
    clasificacion: list[dict[str, Any]] = []
    if cmi_df.empty or "Id" not in cmi_df.columns:
        return {"periodicidad": periodicidad, "clasificacion": clasificacion}

    work = cmi_df.copy()
    work["Id_norm"] = work["Id"].astype(str).str.strip()
    work = work[work["Id_norm"].isin(active_ids)]
    if work.empty:
        return {"periodicidad": periodicidad, "clasificacion": clasificacion}

    per_col = next(
        (c for c in work.columns if "periodicidad" in c.lower() or c == "Frecuencia"), None
    )
    if per_col:
        vc = work[per_col].fillna("Sin dato").astype(str).value_counts()
        periodicidad = [{"label": str(k), "count": int(v)} for k, v in vc.items()]

    clas_col = next((c for c in work.columns if c == "Clasificacion"), None)
    if clas_col:
        vc = work[clas_col].fillna("Sin dato").astype(str).value_counts()
        clasificacion = [{"label": str(k), "count": int(v)} for k, v in vc.items()]

    return {"periodicidad": periodicidad, "clasificacion": clasificacion}


def build_indicadores_summary(df: pd.DataFrame) -> dict[str, int]:
    if df.empty:
        return {
            "total": 0,
            "metricas": 0,
            "sobrecumplimiento": 0,
            "cumplimiento": 0,
            "alerta": 0,
            "peligro": 0,
        }
    unique = latest_per_indicator(df)
    total = len(unique)
    metricas = int(unique["Id"].nunique()) if "Id" in unique.columns else total
    counts = {"sobrecumplimiento": 0, "cumplimiento": 0, "alerta": 0, "peligro": 0}
    if "Nivel de cumplimiento" in unique.columns:
        for nivel, key in [
            ("Sobrecumplimiento", "sobrecumplimiento"),
            ("Cumplimiento", "cumplimiento"),
            ("Alerta", "alerta"),
            ("Peligro", "peligro"),
        ]:
            counts[key] = int((unique["Nivel de cumplimiento"] == nivel).sum())
    return {"total": total, "metricas": metricas, **counts}


def build_variacion_analisis(df: pd.DataFrame, df_prev: pd.DataFrame) -> dict[str, Any]:
    if df.empty:
        return {"mejoraron": [], "empeoraron": [], "top_riesgo_procesos": []}

    mejoraron, empeoraron = compute_trends(df, df_prev)

    col = "Proceso_padre" if "Proceso_padre" in df.columns else "Proceso"
    top_riesgo: list[dict[str, Any]] = []
    if col in df.columns and "Nivel de cumplimiento" in df.columns:
        riesgo_df = df[df["Nivel de cumplimiento"].isin(["Peligro", "Alerta"])]
        if not riesgo_df.empty:
            agg = (
                riesgo_df.groupby(col)
                .agg(
                    n_riesgo=("Nivel de cumplimiento", "count"),
                    cumplimiento=("cumplimiento_pct", "mean"),
                )
                .reset_index()
                .sort_values("n_riesgo", ascending=False)
                .head(8)
            )
            for _, row in agg.iterrows():
                top_riesgo.append(
                    {
                        "proceso": str(row[col]),
                        "n_riesgo": int(row["n_riesgo"]),
                        "cumplimiento": _safe_float(row["cumplimiento"]),
                    }
                )

    return {
        "mejoraron": mejoraron[:10],
        "empeoraron": empeoraron[:10],
        "top_riesgo_procesos": top_riesgo,
    }


def build_comparativa_procesos(
    df: pd.DataFrame, df_prev: pd.DataFrame | None = None
) -> list[dict[str, Any]]:
    col = "Proceso_padre" if "Proceso_padre" in df.columns else "Proceso"
    if df.empty or col not in df.columns:
        return []
    items: list[dict[str, Any]] = []
    for proceso, group in df.groupby(col, dropna=True):
        if not proceso or str(proceso).strip() in ("", "nan"):
            continue
        actual = (
            _safe_float(group["cumplimiento_pct"].mean())
            if "cumplimiento_pct" in group.columns
            else None
        )
        anterior = None
        variacion = None
        if df_prev is not None and not df_prev.empty and col in df_prev.columns:
            prev_g = df_prev[df_prev[col].astype(str) == str(proceso)]
            if not prev_g.empty:
                anterior = _safe_float(prev_g["cumplimiento_pct"].mean())
                if actual is not None and anterior is not None:
                    variacion = round(actual - anterior, 1)
        criticos = 0
        if "Nivel de cumplimiento" in group.columns:
            criticos = int((group["Nivel de cumplimiento"] == "Peligro").sum())
        estado = cumplimiento_estado(actual)
        tipo = (
            str(group["Tipo de proceso"].iloc[0])
            if "Tipo de proceso" in group.columns and len(group)
            else ""
        )
        items.append(
            {
                "proceso": str(proceso),
                "tipo_proceso": _normalize_tipo(tipo),
                "tipo_color": TIPO_PROCESO_COLORS.get(_normalize_tipo(tipo), "#BDBDBD"),
                "cumplimiento": actual,
                "cumplimiento_anterior": anterior,
                "variacion": variacion,
                "n_indicadores": len(group),
                "n_criticos": criticos,
                "color": cumplimiento_semaforo_color(actual),
                "estado": estado["label"],
                "estado_icon": estado["icon"],
                "estado_color": estado["color"],
            }
        )
    items.sort(key=lambda x: x["cumplimiento"] or 0, reverse=True)
    return items


def build_alertas_criticas(df: pd.DataFrame, *, limit: int = 6) -> list[dict[str, Any]]:
    if df.empty or "Nivel de cumplimiento" not in df.columns:
        return []
    peligro = df[df["Nivel de cumplimiento"] == "Peligro"].copy()
    if peligro.empty:
        return []
    if "cumplimiento_pct" in peligro.columns:
        peligro = peligro.sort_values("cumplimiento_pct", ascending=True)
    items: list[dict[str, Any]] = []
    for _, row in peligro.head(limit).iterrows():
        cumpl = _safe_float(row.get("cumplimiento_pct"))
        brecha = round(80 - cumpl, 1) if cumpl is not None else None
        proc = str(row.get("Proceso_padre", row.get("Proceso", "Sin proceso")))
        tipo = _normalize_tipo(str(row.get("Tipo de proceso", "")))
        diag = (
            f"Cumplimiento {cumpl:.1f}% — requiere plan de acción inmediato."
            if cumpl is not None
            else "Sin dato de cumplimiento reportado."
        )
        items.append(
            {
                "id": str(row.get("Id", "")),
                "indicador": str(row.get("Indicador", "Sin nombre")),
                "proceso": proc,
                "tipo_proceso": tipo,
                "tipo_color": TIPO_PROCESO_COLORS.get(tipo, "#BDBDBD"),
                "cumplimiento": cumpl,
                "brecha_pp": brecha,
                "nivel": "Peligro",
                "diagnostico": diag,
            }
        )
    return items


def build_ejecucion_variacion(
    df: pd.DataFrame, df_prev: pd.DataFrame, *, limit: int = 8
) -> dict[str, list[dict[str, Any]]]:
    """Mayor variación de ejecución entre el corte anterior y el actual.

    Los indicadores no comparten unidad ($, %, conteos, días…) ni sentido (mayor/menor es
    mejor), por lo que NO se ordena por la diferencia absoluta de ejecución: se agrupa y ordena
    por la variación del % de cumplimiento en puntos porcentuales (`delta_cumpl_pp`), que sí es
    comparable y ya incorpora meta y sentido. Cada fila trae además el signo/unidad y los
    decimales propios del indicador para formatear ejecución, ejecución previa y delta
    (frontend: fmtValorSigno)."""
    vacio: dict[str, list[dict[str, Any]]] = {"positiva": [], "negativa": []}
    if df.empty or df_prev.empty or "Id" not in df.columns or "Id" not in df_prev.columns:
        return vacio
    if "Ejecucion" not in df.columns or "Ejecucion" not in df_prev.columns:
        return vacio

    extra = [
        c
        for c in ("Meta_Signo", "Ejecucion_s", "Decimales_Ejecucion", "Decimales_Meta")
        if c in df.columns
    ]
    if "cumplimiento_pct" not in df.columns or "cumplimiento_pct" not in df_prev.columns:
        return vacio
    cur = df[
        ["Id", "Indicador", "Ejecucion", "cumplimiento_pct", "Mes", "Anio", *extra]
    ].drop_duplicates(subset=["Id"], keep="last")
    prev_cols = ["Id", "Ejecucion", "cumplimiento_pct"] + [
        c for c in ("Mes", "Anio") if c in df_prev.columns
    ]
    prev = df_prev[prev_cols].drop_duplicates(subset=["Id"], keep="last")
    merged = cur.merge(prev, on="Id", suffixes=("", "_prev"))
    ejec = pd.to_numeric(merged["Ejecucion"], errors="coerce")
    ejec_prev = pd.to_numeric(merged["Ejecucion_prev"], errors="coerce")
    merged["delta"] = ejec - ejec_prev
    merged["delta_cumpl_pp"] = pd.to_numeric(
        merged["cumplimiento_pct"], errors="coerce"
    ) - pd.to_numeric(merged["cumplimiento_pct_prev"], errors="coerce")
    merged = merged.dropna(subset=["delta", "delta_cumpl_pp"])
    merged = merged[(merged["delta"] != 0) & (merged["delta_cumpl_pp"] != 0)]
    if merged.empty:
        return vacio

    def _signo(r: pd.Series) -> str | None:
        for col in ("Ejecucion_s", "Meta_Signo"):
            v = r.get(col)
            if isinstance(v, str) and v.strip() and v.strip().lower() != "nan":
                return v.strip()
        return None

    def _periodo(mes: Any, anio: Any) -> str:
        if pd.isna(mes) or pd.isna(anio):
            return ""
        # Mes llega como número o como nombre ("Junio"), según la fuente.
        try:
            nombre = mes_nombre(int(float(mes)))
        except (TypeError, ValueError):
            nombre = str(mes).strip()
        return f"{nombre} {int(float(anio))}"

    def _row(r: pd.Series) -> dict[str, Any]:
        periodo = _periodo(r.get("Mes"), r.get("Anio"))
        dec = r.get("Decimales_Ejecucion")
        if pd.isna(dec):
            dec = r.get("Decimales_Meta")
        return {
            "indicador": str(r.get("Indicador", r["Id"])),
            "ejecucion": _safe_float(r.get("Ejecucion")),
            "ejecucion_previa": _safe_float(r.get("Ejecucion_prev")),
            "delta": round(float(r["delta"]), 4),
            "delta_cumpl_pp": round(float(r["delta_cumpl_pp"]), 1),
            "signo": _signo(r),
            "decimales": None if pd.isna(dec) else int(float(dec)),
            "periodo": periodo or "—",
            "periodo_previo": _periodo(r.get("Mes_prev"), r.get("Anio_prev")) or None,
        }

    pos = (
        merged[merged["delta_cumpl_pp"] > 0].sort_values("delta_cumpl_pp", ascending=False)
    ).head(limit)
    neg = merged[merged["delta_cumpl_pp"] < 0].sort_values("delta_cumpl_pp").head(limit)
    return {
        "positiva": [_row(r) for _, r in pos.iterrows()],
        "negativa": [_row(r) for _, r in neg.iterrows()],
    }


def build_vista_global(
    df_global: pd.DataFrame,
    df_base_year: pd.DataFrame,
    cmi_catalog: pd.DataFrame,
    *,
    anio: int,
    mes_corte: int,
    base_year: int,
    base_mes: int | None,
) -> dict[str, Any]:
    latest = latest_per_indicator(df_global)
    active_ids = (
        set(latest["Id"].astype(str).str.strip().tolist()) if "Id" in latest.columns else set()
    )
    cumpl_global = avg_cumplimiento(latest)
    cumpl_base = avg_cumplimiento(df_base_year)
    return {
        "mes_corte": mes_corte,
        "mes_nombre": mes_nombre(mes_corte),
        "kpis": build_procesos_kpis(latest),
        "banner": build_banner(
            latest,
            anio=anio,
            mes=mes_corte,
            base_year=base_year,
            base_month=base_mes,
            cumpl_global=cumpl_global,
            cumpl_base=cumpl_base,
        ),
        "distribucion_nivel": build_distribucion_nivel(latest),
        "tipo_proceso_cards": build_tipo_proceso_cards(latest, df_base_year),
        "proceso_bars": build_proceso_bars(latest, df_base_year),
        "catalog_charts": build_catalog_charts(cmi_catalog, active_ids),
        "procesos_detalle": build_procesos_detalle(latest),
        "unidades_detalle": build_unidades_detalle(latest),
        "comparativa_procesos": build_comparativa_procesos(latest, df_base_year),
        "variacion": build_variacion_analisis(
            latest, df_base_year if not df_base_year.empty else df_global
        ),
        "alertas_criticas": build_alertas_criticas(latest),
        "brecha_ambiental": build_brecha_ambiental(latest),
    }


def build_filtros_options(
    tracking: pd.DataFrame,
    map_df: pd.DataFrame,
    cmi_df: pd.DataFrame,
    *,
    anio: int,
    mes: int | None = None,
    jerarquia: pd.DataFrame | None = None,
    por_cierre: bool = False,
) -> dict[str, Any]:
    # Debe reflejar el mismo tope que cmi_service.MAX_ANIO_FILTROS.
    anios = (
        sorted(
            a
            for a in pd.to_numeric(tracking["Anio"], errors="coerce")
            .dropna()
            .astype(int)
            .unique()
            .tolist()
            if a <= 2026
        )
        if "Anio" in tracking.columns
        else [anio]
    )

    prepared = prepare_tracking(tracking, map_df, jerarquia)
    # Usa el mes seleccionado por el usuario cuando se conoce — si las opciones
    # siempre se calculan sobre default_mes() (el mes más reciente con datos),
    # un corte reciente aún incompleto (ej. Diciembre con solo 6 de ~34 procesos
    # reportados) hace que el filtro de Proceso muestre solo esos 6, incluso
    # cuando el usuario tiene seleccionado un mes anterior ya cerrado (ej. Junio)
    # con los ~34 procesos completos. Ver bugfix "filtros 2026 Junio".
    if por_cierre:
        # Un solo corte por año: el último reporte de cada indicador (no hay filtro por mes).
        year_slice = ultimo_registro_anual(
            prepared[pd.to_numeric(prepared["Anio"], errors="coerce") == int(anio)]
            if "Anio" in prepared.columns
            else prepared,
            mes_corte=12,
        )
    else:
        mes_eff = mes if mes is not None else default_mes(tracking, anio)
        year_slice = filter_by_anio_mes(prepared, anio=anio, mes=mes_eff)

    unidades = (
        sorted(year_slice["Unidad"].replace("", pd.NA).dropna().astype(str).unique().tolist())
        if "Unidad" in year_slice.columns
        else []
    )
    procesos = (
        sorted(year_slice["Proceso_padre"].dropna().astype(str).unique().tolist())
        if "Proceso_padre" in year_slice.columns
        else []
    )
    subprocesos = (
        sorted(year_slice["Subproceso_final"].dropna().astype(str).unique().tolist())
        if "Subproceso_final" in year_slice.columns
        else []
    )

    subprocesos_por_proceso: dict[str, list[str]] = {}
    if "Proceso_padre" in year_slice.columns and "Subproceso_final" in year_slice.columns:
        for proc in procesos:
            subs = sorted(
                year_slice[year_slice["Proceso_padre"] == proc]["Subproceso_final"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )
            subprocesos_por_proceso[proc] = subs

    clasificaciones: list[str] = []
    if "Clasificacion" in year_slice.columns:
        clasificaciones = sorted(year_slice["Clasificacion"].dropna().astype(str).unique().tolist())

    frecuencias: list[str] = []
    freq_col = next((c for c in ["Frecuencia", "Periodicidad"] if c in year_slice.columns), None)
    if freq_col:
        frecuencias = sorted(year_slice[freq_col].dropna().astype(str).unique().tolist())

    meses_disp: list[int] = []
    if por_cierre:
        pass
    elif "Anio" in tracking.columns and "Mes" in tracking.columns:
        sub = tracking[pd.to_numeric(tracking["Anio"], errors="coerce") == int(anio)]
        nums = sub["Mes"].apply(mes_to_num).dropna().astype(int).unique().tolist()
        meses_disp = sorted(nums)

    return {
        "anios": anios,
        "meses": [] if por_cierre else (meses_disp or list(range(1, 13))),
        "mes_default": ultimo_mes_reportado(year_slice) or 12
        if por_cierre
        else default_mes(tracking, anio),
        "unidades": unidades,
        "procesos": procesos,
        "subprocesos": subprocesos,
        "subprocesos_por_proceso": subprocesos_por_proceso,
        "clasificaciones": clasificaciones,
        "frecuencias": frecuencias,
    }


def etiqueta_reporte(fecha: object) -> str | None:
    f = pd.to_datetime(fecha, errors="coerce")
    return None if pd.isna(f) else f"{mes_nombre(int(f.month))[:3]} {f.year}"


def build_indicadores_procesos_listado(df: pd.DataFrame) -> list[dict[str, Any]]:
    if not df.empty and "Fecha" in df.columns:
        df = df.copy()
        df["ultimo_reporte"] = df["Fecha"].map(etiqueta_reporte)
        if "Estado_Cierre" in df.columns:
            df["estado_cierre"] = df["Estado_Cierre"]
        else:
            df["estado_cierre"] = df["Fecha"].map(estado_cierre_fila)
    cols_extra = [
        c
        for c in ["Proceso_padre", "Subproceso_final", "Unidad", "Tipo de proceso"]
        if c in df.columns
    ]
    base_cols = [
        c
        for c in [
            "Id",
            "Indicador",
            "Proceso",
            "Subproceso",
            "Meta",
            "Ejecucion",
            "Meta_Signo",
            "Ejecucion_s",
            "EjecS",
            "Decimales_Meta",
            "Decimales_Ejecucion",
            "cumplimiento_pct",
            "Nivel de cumplimiento",
            "Nivel sobreejecucion",
            "Clasificacion",
            "Frecuencia",
            "Periodicidad",
            "Anio",
            "Mes",
            "ultimo_reporte",
            "estado_cierre",
            "TipoCalculo",
        ]
        if c in df.columns
    ]
    return df_to_records(df, list(dict.fromkeys(base_cols + cols_extra)))


def avg_cumplimiento(df: pd.DataFrame) -> float | None:
    if df.empty or "cumplimiento_pct" not in df.columns:
        return None
    val = _safe_float(df["cumplimiento_pct"].mean())
    return round(val, 1) if val is not None else None


def default_anio_procesos(anios: list[int]) -> int:
    if 2026 in anios:
        return 2026
    if 2025 in anios:
        return 2025
    if anios:
        return anios[-1]
    return date.today().year


def process_variation_for_rpp(
    df_current: pd.DataFrame,
    df_prev: pd.DataFrame,
    display_col: str = "Proceso_padre",
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if df_current.empty or df_prev.empty or display_col not in df_current.columns:
        return [], []
    curr = (
        df_current.groupby(display_col, dropna=False)
        .agg(actual=("cumplimiento_pct", "mean"))
        .reset_index()
    )
    prev = (
        df_prev.groupby(display_col, dropna=False)
        .agg(prev=("cumplimiento_pct", "mean"))
        .reset_index()
    )
    merged = curr.merge(prev, on=display_col, how="left")
    merged["change"] = merged["actual"] - pd.to_numeric(merged["prev"], errors="coerce")
    merged = merged.dropna(subset=["change"])
    if merged.empty:
        return [], []
    best = [
        {"name": str(r[display_col]), "change": round(float(r["change"]), 1)}
        for _, r in merged.sort_values("change", ascending=False).head(5).iterrows()
    ]
    worst = [
        {"name": str(r[display_col]), "change": round(float(r["change"]), 1)}
        for _, r in merged.sort_values("change", ascending=True).head(5).iterrows()
    ]
    return best, worst


def build_propuesta_accion(df: pd.DataFrame, proceso: str = "Todos") -> dict[str, Any]:
    if df.empty:
        return {
            "proceso": proceso,
            "plan_mejoramiento": "Sin datos para priorizar acciones.",
            "pdi": "Definir metas e hitos por indicador.",
            "sga": "Validar integración de evidencias y trazabilidad.",
            "retos": "Completar reporte de indicadores faltantes.",
            "top_criticos": [],
        }
    work = df.copy()
    if "cumplimiento_pct" not in work.columns:
        return build_propuesta_accion(pd.DataFrame(), proceso)
    riesgos = work[work["cumplimiento_pct"].notna() & (work["cumplimiento_pct"] < 80)]
    top = riesgos.nsmallest(3, "cumplimiento_pct") if not riesgos.empty else pd.DataFrame()
    top_criticos = []
    for _, r in top.iterrows():
        top_criticos.append(
            {
                "indicador": str(r.get("Indicador", "Sin nombre")),
                "proceso": str(r.get("Proceso_padre", r.get("Proceso", "Sin proceso"))),
                "cumplimiento": _safe_float(r.get("cumplimiento_pct")),
            }
        )
    nombres = ", ".join(t["indicador"] for t in top_criticos) or "Ninguno crítico en el corte"
    return {
        "proceso": proceso,
        "plan_mejoramiento": f"Priorizar cierre de brechas en: {nombres}.",
        "pdi": "Alinear metas por resultados históricos y capacidad operativa.",
        "sga": "Fortalecer controles de calidad de dato y consistencia de soportes.",
        "retos": "Sostener cumplimiento mensual y reducir dispersión entre subprocesos.",
        "top_criticos": top_criticos,
    }


def generate_ficha_narrativa_heuristica(
    *,
    nombre: str,
    meta: Any,
    ejecucion: Any,
    nivel: str,
    cumplimiento: float | None,
    proceso: str | None = None,
) -> dict[str, str]:
    cump = cumplimiento or 0.0
    if cump >= 100:
        estado = "El indicador supera la meta y presenta desempeño favorable en el corte actual."
        riesgo = "Riesgo de sobreconfianza: sostener el resultado sin plan de continuidad puede generar retroceso."
        accion = "Estandarizar las prácticas que explican el resultado y documentar un plan de sostenibilidad."
    elif cump >= 95:
        estado = "El indicador está cerca de la meta y requiere ajuste fino para consolidar cumplimiento."
        riesgo = (
            "Riesgo de cierre en alerta por variaciones menores de ejecución o retrasos operativos."
        )
        accion = "Definir acciones de corto plazo con responsables y seguimiento semanal hasta el próximo corte."
    else:
        estado = (
            "El indicador presenta brecha frente a la meta y requiere intervención prioritaria."
        )
        riesgo = (
            "Riesgo de incumplimiento del objetivo asociado y efectos en el balance del proceso."
        )
        accion = (
            "Implementar plan de recuperación con metas parciales y control quincenal de avance."
        )

    proc_ctx = f" en el proceso <strong>{proceso}</strong>" if proceso else ""
    texto = (
        f"<strong>Diagnóstico:</strong> {estado}{proc_ctx}<br/>"
        f"<strong>Riesgo principal:</strong> {riesgo}<br/>"
        f"<strong>Recomendación táctica:</strong> {accion}<br/>"
        f"<strong>Meta / Ejecución:</strong> {meta} / {ejecucion} — Nivel: {nivel}"
    )
    return {
        "diagnostico": estado,
        "riesgo": riesgo,
        "recomendacion": accion,
        "texto_html": texto,
        "fuente": "heuristica",
    }


def generate_proceso_narrativa_heuristica(
    proceso: str,
    df: pd.DataFrame,
    *,
    cumplimiento_promedio: float | None = None,
) -> dict[str, Any]:
    total = len(df)
    riesgo = 0
    if "Nivel de cumplimiento" in df.columns:
        riesgo = int(df["Nivel de cumplimiento"].isin(["Peligro", "Alerta"]).sum())
    cump = (
        cumplimiento_promedio
        if cumplimiento_promedio is not None
        else _safe_float(df["cumplimiento_pct"].mean())
    )
    cump = cump or 0.0
    riesgo_ratio = (riesgo / total) if total else 0.0

    if cump >= 100:
        estado = f"El proceso <strong>{proceso}</strong> supera la meta institucional en el corte actual."
        color = "#16A34A"
    elif cump >= 95:
        estado = f"El proceso <strong>{proceso}</strong> mantiene desempeño estable con brechas acotadas."
        color = "#2563EB"
    else:
        estado = f"El proceso <strong>{proceso}</strong> presenta desviación relevante y requiere priorización."
        color = "#DC2626"

    dir_1 = "Enfocar seguimiento semanal en indicadores en riesgo del proceso."
    dir_2 = (
        "Activar comité táctico con responsables por subproceso."
        if riesgo_ratio > 0.25
        else "Replicar buenas prácticas de subprocesos con mejor desempeño."
    )
    texto = (
        f"{estado} Hay <strong>{riesgo}</strong> indicadores en alerta/peligro sobre <strong>{total}</strong>.<br/>"
        f"<strong>Directriz 1:</strong> {dir_1}<br/>"
        f"<strong>Directriz 2:</strong> {dir_2}"
    )
    return {
        "titulo": "Análisis heurístico del proceso",
        "estado_color": color,
        "foco_urgente": dir_1,
        "directrices": [dir_1, dir_2],
        "texto_html": texto,
        "fuente": "heuristica",
    }


def build_historico_catalog(
    df: pd.DataFrame, tracking: pd.DataFrame, *, limit: int = 50
) -> list[dict[str, Any]]:
    if df.empty or tracking.empty or "Id" not in df.columns:
        return []
    ids = set(df["Id"].astype(str).str.strip().tolist())
    hist = tracking[tracking["Id"].astype(str).str.strip().isin(ids)].copy()
    if hist.empty:
        return []

    catalog: list[dict[str, Any]] = []
    name_col = "Indicador" if "Indicador" in hist.columns else "Id"
    for ind_id, group in hist.groupby("Id"):
        ind_id = str(ind_id)
        name = str(group[name_col].iloc[0]) if name_col in group.columns else ind_id
        puntos = []
        sort_cols = [c for c in ["Anio", "Mes"] if c in group.columns]
        sorted_g = group.sort_values(sort_cols) if sort_cols else group
        for _, row in sorted_g.iterrows():
            mes_num = mes_to_num(row.get("Mes"))
            anio = row.get("Anio")
            periodo = (
                f"{anio}-{int(mes_num):02d}" if mes_num and pd.notna(anio) else str(anio or "")
            )
            pct = row.get("cumplimiento_pct") or row.get("Cumplimiento_norm")
            if pct is not None:
                try:
                    val = float(pct)
                    if val <= 2:
                        val *= 100
                except (TypeError, ValueError):
                    val = None
            else:
                val = None
            if val is not None:
                puntos.append({"periodo": periodo, "cumplimiento": round(val, 1)})
        if puntos:
            catalog.append({"id": ind_id, "indicador": name, "puntos": puntos})
        if len(catalog) >= limit:
            break
    return catalog


_EXPORT_COLS = {
    "Id": "Código",
    "Indicador": "Indicador",
    "Proceso_padre": "Proceso",
    "Subproceso_final": "Subproceso",
    "Unidad": "Unidad",
    "Tipo de proceso": "Tipo de proceso",
    "Meta": "Meta",
    "Ejecucion": "Ejecución",
    "cumplimiento_pct": "Cumplimiento (%)",
    "Nivel de cumplimiento": "Estado",
    "Nivel sobreejecucion": "Sobreejecución",
    "Clasificacion": "Clasificación",
    "Frecuencia": "Frecuencia",
    "Periodicidad": "Periodicidad",
    "Anio": "Año",
    "Mes": "Mes",
}


def build_export_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    present = {k: v for k, v in _EXPORT_COLS.items() if k in df.columns}
    if not present:
        return df.copy()
    out = df[list(present.keys())].copy().rename(columns=present)
    if "Cumplimiento (%)" in out.columns:
        out["Cumplimiento (%)"] = pd.to_numeric(out["Cumplimiento (%)"], errors="coerce").round(1)
    return out


def build_export_excel_bytes(df: pd.DataFrame) -> bytes:
    import io

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    df_exp = build_export_dataframe(df)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_exp.to_excel(writer, sheet_name="Indicadores CMI Procesos", index=False)
        ws = writer.sheets["Indicadores CMI Procesos"]
        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr_fill = PatternFill(start_color="1A3A5C", end_color="1A3A5C", fill_type="solid")
        hdr_font = Font(color="FFFFFF", bold=True, size=11)
        for col_idx in range(1, len(df_exp.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        col_widths = {"Indicador": 48, "Proceso": 28, "Subproceso": 28}
        for col_idx, col_name in enumerate(df_exp.columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 18)
        ws.freeze_panes = "A2"
    return buf.getvalue()


def build_analisis_avanzado(
    df: pd.DataFrame,
    df_prev: pd.DataFrame,
    df_base_year: pd.DataFrame,
    tracking: pd.DataFrame,
    *,
    proceso: str = "Todos",
    base_anio: int,
) -> dict[str, Any]:
    col = "Proceso_padre" if "Proceso_padre" in df.columns else "Proceso"
    # No sustituir df (filtrado) por un dataset global cuando está vacío: comparar
    # el corte filtrado contra un df_prev/df_base_year sin filtrar produce falsos
    # "mejora"/"riesgo" de procesos que ni siquiera pertenecen al filtro aplicado.
    df_comp_prev = df_prev if not df_prev.empty else df_base_year
    best_proc, worst_proc = (
        process_variation_for_rpp(df, df_comp_prev, display_col=col) if not df.empty else ([], [])
    )
    propuesta = build_propuesta_accion(df, proceso=proceso)
    narrativa = generate_proceso_narrativa_heuristica(
        proceso if proceso != "Todos" else "Todos los procesos",
        df,
    )
    variacion = build_variacion_analisis(df, df_prev if not df_prev.empty else df_base_year)
    historico = build_historico_catalog(df, tracking)

    mejora_txt = (
        f"{best_proc[0]['name']} registra la mayor variación positiva: {best_proc[0]['change']:+.1f}% vs {base_anio}."
        if best_proc
        else "Sin datos comparativos de mejora."
    )
    riesgo_txt = (
        f"{worst_proc[0]['name']} presenta la mayor caída: {worst_proc[0]['change']:+.1f}% vs {base_anio}."
        if worst_proc
        else "Sin datos comparativos de riesgo."
    )

    return {
        "propuesta_accion": propuesta,
        "variacion_procesos": {"mejoraron": best_proc, "empeoraron": worst_proc},
        "insights": {"mejora_proceso": mejora_txt, "riesgo_proceso": riesgo_txt},
        "narrativa_proceso": narrativa,
        "variacion_indicadores": variacion,
        "historico_indicadores": historico,
    }


# ── Mapa de procesos: Proceso → Subproceso → Indicador, histórico anual ──────────────


def ultimo_registro_anual(year_df: pd.DataFrame, *, mes_corte: int) -> pd.DataFrame:
    """Último registro por Id del año hasta `mes_corte` (inclusive).

    El valor anual de cada indicador ya viene agregado por el ETL según su TipoCalculo
    (Cierre / Promedio / Acumulado); aquí solo se elige el último periodo reportado.
    """
    if year_df.empty or "Id" not in year_df.columns:
        return year_df.iloc[0:0]
    out = year_df.copy()
    out["_mes_num"] = out["Mes"].apply(mes_to_num) if "Mes" in out.columns else 12.0
    out = out[out["_mes_num"].fillna(12) <= mes_corte].sort_values("_mes_num", kind="stable")
    return out.drop_duplicates(subset=["Id"], keep="last").drop(columns=["_mes_num"])


def ultimo_mes_reportado(df: pd.DataFrame) -> int | None:
    """Mes más reciente entre los últimos reportes de los indicadores del corte."""
    if df.empty or "Mes" not in df.columns:
        return None
    meses = df["Mes"].apply(mes_to_num).dropna()
    return int(meses.max()) if not meses.empty else None


def estado_cierre_fila(fecha: object, anio_curso: int | None = None) -> str:
    """'Parcial' si el corte es del año en curso; 'Final' en cualquier año anterior.
    Misma regla que `scripts/etl/purga.py::estado_cierre` (columna Estado_Cierre del ETL)."""
    curso = date.today().year if anio_curso is None else anio_curso
    f = pd.to_datetime(fecha, errors="coerce")
    if pd.isna(f):
        return ""
    return "Parcial" if f.year >= curso else "Final"


def _serie_anual(frames: dict[int, pd.DataFrame], anios: list[int]) -> dict[str, float | None]:
    return {str(a): avg_cumplimiento(frames[a]) if a in frames else None for a in anios}


def _variaciones(serie: dict[str, float | None], anios: list[int]) -> dict[str, float | None]:
    """Variación en puntos porcentuales entre años consecutivos de los últimos tres."""
    out: dict[str, float | None] = {}
    for a, b in zip(anios[-3:-1], anios[-2:], strict=False):
        va, vb = serie.get(str(a)), serie.get(str(b))
        out[f"{a}-{b}"] = round(vb - va, 1) if va is not None and vb is not None else None
    return out


def _conteo_riesgo(df: pd.DataFrame) -> int:
    if df.empty or "Nivel de cumplimiento" not in df.columns:
        return 0
    return int(df["Nivel de cumplimiento"].isin(["Peligro", "Alerta"]).sum())


def build_mapa_procesos(
    frames: dict[int, pd.DataFrame],
    jerarquia: pd.DataFrame,
    *,
    anios: list[int],
    anio_parcial: int | None,
) -> dict[str, Any]:
    """Árbol Proceso → Subproceso → Indicador con cumplimiento de cada año.

    `frames[año]` = último registro por indicador de ese año (ver `ultimo_registro_anual`),
    ya con Proceso_padre/Subproceso_final asignados por Id. El promedio de subproceso y
    proceso se calcula sobre los cumplimientos de sus indicadores.
    """
    tipo_por_id: dict[str, str] = {}
    if jerarquia is not None and not jerarquia.empty:
        tipo_por_id = dict(zip(jerarquia["Id"], jerarquia["TipoCalculo"], strict=False))

    def _tipo_calculo(id_val: object) -> str:
        s = _id_base(id_val)
        return tipo_por_id.get(s) or tipo_por_id.get(s.split(".")[0], "") or ""

    todo = [f.assign(_anio=a) for a, f in frames.items() if not f.empty]
    if not todo:
        return {
            "anios": anios,
            "anio_parcial": anio_parcial,
            "procesos": [],
            "mix_niveles": {},
            "sin_tipo_calculo": 0,
        }
    big = pd.concat(todo, ignore_index=True)
    sin_tipo: set[str] = set()
    procesos: list[dict[str, Any]] = []

    for proceso, gp in big.groupby("Proceso_padre", dropna=True, sort=True):
        if not str(proceso).strip():
            continue
        gp_frames = {a: g for a, g in gp.groupby("_anio")}
        serie_p = _serie_anual(gp_frames, anios)
        subprocesos: list[dict[str, Any]] = []
        for sub, gs in gp.groupby("Subproceso_final", dropna=True, sort=True):
            gs_frames = {a: g for a, g in gs.groupby("_anio")}
            serie_s = _serie_anual(gs_frames, anios)
            inds: list[dict[str, Any]] = []
            for id_val, gi in gs.groupby("Id", sort=False):
                gi_frames = {a: g for a, g in gi.groupby("_anio")}
                actual = gi.sort_values("_anio").iloc[-1]
                tc = _tipo_calculo(id_val)
                if not tc:
                    sin_tipo.add(_id_base(id_val))
                fecha = pd.to_datetime(actual.get("Fecha"), errors="coerce")
                estado = str(actual.get("Estado_Cierre", "") or "").strip() or estado_cierre_fila(fecha)
                inds.append(
                    {
                        "id": _id_base(id_val),
                        "indicador": str(actual.get("Indicador", "")),
                        "tipo_calculo": tc or None,
                        "periodicidad": str(actual.get("Periodicidad", "") or "") or None,
                        "ultimo_reporte": None
                        if pd.isna(fecha)
                        else f"{mes_nombre(int(fecha.month))[:3]} {fecha.year}",
                        "estado_cierre": estado or None,
                        "anios": _serie_anual(gi_frames, anios),
                        "nivel": str(actual.get("Nivel de cumplimiento", "") or "") or None,
                    }
                )
            inds.sort(key=lambda x: (x["indicador"], x["id"]))
            ultimo = max(gs_frames) if gs_frames else None
            subprocesos.append(
                {
                    "subproceso": str(sub),
                    "anios": serie_s,
                    "variaciones": _variaciones(serie_s, anios),
                    "n_indicadores": len(inds),
                    "en_riesgo": _conteo_riesgo(gs_frames[ultimo]) if ultimo else 0,
                    "indicadores": inds,
                }
            )
        ultimo_p = max(gp_frames) if gp_frames else None
        procesos.append(
            {
                "proceso": str(proceso),
                "unidad": str(gp["Unidad"].iloc[0]) if "Unidad" in gp.columns else "",
                "tipo_proceso": str(gp["Tipo de proceso"].iloc[0])
                if "Tipo de proceso" in gp.columns
                else "",
                "anios": serie_p,
                "variaciones": _variaciones(serie_p, anios),
                "n_indicadores": int(gp["Id"].nunique()),
                "en_riesgo": _conteo_riesgo(gp_frames[ultimo_p]) if ultimo_p else 0,
                "subprocesos": subprocesos,
            }
        )
    procesos.sort(key=lambda p: (p["tipo_proceso"], p["proceso"]))
    mix: dict[str, dict[str, int]] = {}
    for a, f in frames.items():
        if f.empty or "Nivel de cumplimiento" not in f.columns:
            continue
        mix[str(a)] = {
            str(k): int(v) for k, v in f.drop_duplicates("Id")["Nivel de cumplimiento"].value_counts().items()
        }
    return {
        "anios": anios,
        "anio_parcial": anio_parcial,
        "procesos": procesos,
        "mix_niveles": mix,
        "sin_tipo_calculo": len(sin_tipo),
    }


# ── Tendencia mensual del año en curso (visual con filtro de mes) ─────────────────────


def build_tendencia_mensual(
    year_df: pd.DataFrame,
    *,
    anio: int,
    ultimo_cerrado: int,
    proceso: str | None = None,
    subproceso: str | None = None,
    mes_sel: int | None = None,
) -> dict[str, Any]:
    """Cumplimiento de los 12 meses del año en curso.

    Cada mes M toma, por indicador, su último corte acumulado hasta M (Consolidado Mensual: ya
    aplica Cierre / Promedio / Acumulado). Los meses posteriores al último mes cerrado se
    devuelven vacíos. `mes_sel` (solo este filtro, ≤ último mes cerrado) elige el corte del que
    se detalla el cumplimiento por proceso o subproceso.
    """
    base = apply_ui_filters(year_df, proceso=proceso, subproceso=subproceso)
    serie: list[dict[str, Any]] = []
    for m in range(1, 13):
        item: dict[str, Any] = {
            "mes": m,
            "nombre": mes_nombre(m),
            "cerrado": m <= ultimo_cerrado,
            "cumplimiento": None,
            "n_indicadores": 0,
            "n_reportes_mes": 0,
            "en_riesgo": 0,
        }
        if m <= ultimo_cerrado:
            corte = ultimo_registro_anual(base, mes_corte=m)
            if not corte.empty:
                meses = corte["Mes"].apply(mes_to_num)
                item.update(
                    cumplimiento=avg_cumplimiento(corte),
                    n_indicadores=int(corte["Id"].nunique()),
                    n_reportes_mes=int((meses == m).sum()),
                    en_riesgo=_conteo_riesgo(corte),
                )
        serie.append(item)

    sel = min(max(int(mes_sel or ultimo_cerrado), 1), max(ultimo_cerrado, 1))
    corte_sel = ultimo_registro_anual(base, mes_corte=sel) if ultimo_cerrado >= 1 else base.iloc[0:0]
    col = "Subproceso_final" if (proceso and proceso != "Todos") else "Proceso_padre"
    detalle: list[dict[str, Any]] = []
    if not corte_sel.empty and col in corte_sel.columns:
        for nombre, g in corte_sel.groupby(col, dropna=True):
            if not str(nombre).strip():
                continue
            detalle.append(
                {
                    "nombre": str(nombre),
                    "cumplimiento": avg_cumplimiento(g),
                    "n_indicadores": int(g["Id"].nunique()),
                    "en_riesgo": _conteo_riesgo(g),
                }
            )
        detalle.sort(key=lambda d: (d["cumplimiento"] is None, d["cumplimiento"] or 0))
    return {
        "anio": anio,
        "ultimo_mes_cerrado": ultimo_cerrado,
        "mes_seleccionado": sel,
        "mes_seleccionado_nombre": mes_nombre(sel),
        "agrupado_por": "subproceso" if col == "Subproceso_final" else "proceso",
        "serie": serie,
        "detalle": detalle,
    }
