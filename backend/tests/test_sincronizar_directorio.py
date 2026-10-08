"""scripts/sincronizar_directorio_indicadores.py: detecta indicadores nuevos de Kawak/API,
los agrega al catálogo y a la hoja de cada PDI activo, y solicita su asignación estratégica."""

import importlib
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
CATALOGO_REAL = ROOT / "data" / "raw" / "Catalogo de Indicadores.xlsx"

sys.path.insert(0, str(ROOT / "scripts"))
sync = importlib.import_module("sincronizar_directorio_indicadores")


def _fuente(nombre: str) -> dict:
    return {"Indicador": nombre, "Proceso": "DATOS Y ANALÍTICA", "Clasificacion": "Estratégico",
            "Periodicidad": "Semestral", "Sentido": "Positivo", "Tipo_API": "Eficacia", "Estado": "Activo"}


def test_detectar_nuevos_solo_lo_que_el_catalogo_no_tiene():
    kawak = {"1": _fuente("Uno"), "9001": _fuente("Nuevo Kawak")}
    api = {"1": _fuente("Uno"), "9002": _fuente("Nuevo API")}
    nuevos = sync.detectar_nuevos({"1"}, kawak, api)
    assert list(nuevos) == ["9001", "9002"]
    assert nuevos["9001"]["Indicador"] == "Nuevo Kawak"


@pytest.fixture
def catalogo(tmp_path, monkeypatch):
    if not CATALOGO_REAL.exists():
        pytest.skip("catálogo real no disponible")
    copia = tmp_path / "cat.xlsx"
    shutil.copy2(CATALOGO_REAL, copia)
    # Lo que ya esté pendiente en el catálogo real no debe interferir con estas pruebas
    wb = openpyxl.load_workbook(copia)
    ws = wb["PDI_2026_2030"]
    cols = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        if r[cols.index("Origen")].value == sync.ORIGEN_NUEVO and r[cols.index("Linea")].value in (None, ""):
            r[cols.index("Linea")].value = "x"  # ya "asignado": que no interfiera
    wb.save(copia)
    monkeypatch.setattr(sync, "leer_kawak", lambda: ({"9001": _fuente("Nuevo A"), "9002": _fuente("Nuevo B")}, set()))
    monkeypatch.setattr(sync, "leer_api", lambda: {})
    monkeypatch.setattr(sync, "BACKUPS_DIR", tmp_path / "bak")
    return copia


def _fila(wb, hoja: str, ids: str) -> dict:
    ws = wb[hoja]
    enc = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if str(r[0]) == ids:
            return dict(zip(enc, r, strict=False))
    return {}


def test_sin_terminal_agrega_y_deja_pendiente(catalogo, monkeypatch, capsys):
    monkeypatch.setattr(sync, "hay_terminal", lambda: False)
    assert sync.main(["--catalogo", str(catalogo)]) == 0
    wb = openpyxl.load_workbook(catalogo)
    assert _fila(wb, "Catalogo Indicadores", "9001")["Fuente"] == sync.FUENTE_NUEVO
    fila = _fila(wb, "PDI_2026_2030", "9001")
    assert fila["Origen"] == sync.ORIGEN_NUEVO and fila["PDI"] == 0 and fila["Linea"] is None
    aviso = next(l for l in capsys.readouterr().out.splitlines() if l.startswith("[AVISO]") and "asignación" in l)
    assert "9001" in aviso and "9002" in aviso
    # El PDI cerrado no se toca
    assert _fila(wb, "PDI_2022_2026", "9001") == {}
    # Idempotente: una segunda corrida no duplica filas
    sync.main(["--catalogo", str(catalogo)])
    ws = openpyxl.load_workbook(catalogo)["PDI_2026_2030"]
    assert sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if str(r[0]) == "9001") == 1


def test_con_terminal_solicita_la_asignacion_y_respeta_la_regla_de_meta(catalogo, monkeypatch, capsys):
    monkeypatch.setattr(sync, "hay_terminal", lambda: True)
    # 9001: estratégico (1) → línea 1 → objetivo 1 → meta 1.  9002: Enter = omitir.
    respuestas = iter(["1", "1", "1", "1", ""])
    monkeypatch.setattr("builtins.input", lambda *_: next(respuestas))
    assert sync.main(["--catalogo", str(catalogo)]) == 0
    wb = openpyxl.load_workbook(catalogo)
    a = _fila(wb, "PDI_2026_2030", "9001")
    assert a["PDI"] == 1 and a["Linea"] and a["Objetivo"] and a["Meta"]
    b = _fila(wb, "PDI_2026_2030", "9002")
    assert b["PDI"] == 0 and b["Linea"] is None  # por defecto 0 (de proceso)
    out = capsys.readouterr().out
    aviso = next(l for l in out.splitlines() if l.startswith("[AVISO]") and "asignación" in l)
    assert "9002" in aviso and "9001" not in aviso  # solo queda pendiente el omitido
    assert "Regla de meta" not in out


def test_estrategico_sin_meta_se_reporta(catalogo, monkeypatch, capsys):
    monkeypatch.setattr(sync, "hay_terminal", lambda: True)
    # 9001: PDI=1, línea 1, objetivo 1, meta omitida; luego q.
    respuestas = iter(["1", "1", "1", "", "q"])
    monkeypatch.setattr("builtins.input", lambda *_: next(respuestas))
    sync.main(["--catalogo", str(catalogo)])
    assert "Regla de meta: PDI-2026-2030 · 9001" in capsys.readouterr().out
