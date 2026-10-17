"""ETL: cierres parciales del año en curso (reemplazo por el corte más nuevo) y Estado_Cierre."""

import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from scripts.etl.catalogo import _leer_hoja_catalogo, _tipo_calculo_valido  # noqa: E402
from scripts.etl.purga import (  # noqa: E402
    estado_cierre,
    limpiar_cierres_existentes,
    marcar_estado_cierre,
    resaltar_duplicados_id_anio,
    reemplazar_cierres_parciales,
)

CURSO = 2026
HOY = dt.date(2026, 10, 17)
HEADER = ["Id", "Indicador", "Proceso", "Periodicidad", "Sentido", "Fecha", "Año", "Mes", "Periodo"]


def _hoja(filas):
    ws = openpyxl.Workbook().active
    ws.append(HEADER)
    for id_val, fecha in filas:
        ws.append([id_val, "x", "p", "Mensual", "Positivo", fecha, None, None, None])
    return ws


def _contenido(ws):
    return [(ws.cell(r, 1).value, ws.cell(r, 6).value) for r in range(2, ws.max_row + 1)]


def test_estado_cierre_final_en_todo_año_cerrado_y_parcial_solo_en_el_en_curso():
    assert estado_cierre("2022-06-30", CURSO) == "Final"
    assert estado_cierre("2025-12-31", CURSO) == "Final"
    assert estado_cierre("2026-07-31", CURSO) == "Parcial"
    assert estado_cierre("2026-12-31", CURSO) == "Parcial"
    assert estado_cierre("2027-03-31", 2027) == "Parcial"
    assert estado_cierre("2026-12-31", 2027) == "Final"  # al cambiar el año, el anterior cierra
    assert estado_cierre(None, CURSO) == ""


def test_reemplaza_el_parcial_si_hay_corte_nuevo_y_mantiene_si_no():
    ws = _hoja(
        [
            (217, dt.date(2026, 6, 30)),  # parcial previo: llega un corte más nuevo
            (300, dt.date(2026, 6, 30)),  # sin registro nuevo: se mantiene
            (400, dt.date(2026, 7, 31)),  # el nuevo es más viejo: se mantiene el existente
        ]
    )
    nuevos = [
        {"Id": 217, "fecha": dt.date(2026, 7, 31)},
        {"Id": 400, "fecha": dt.date(2026, 6, 30)},
        {"Id": 500, "fecha": dt.date(2026, 7, 31)},
    ]
    a_escribir, reemplazadas, descartados = reemplazar_cierres_parciales(ws, nuevos, CURSO, HOY)
    assert [(r["Id"], r["fecha"]) for r in a_escribir] == [
        (217, dt.date(2026, 7, 31)),
        (500, dt.date(2026, 7, 31)),
    ]
    assert (reemplazadas, descartados) == (1, 1)
    assert _contenido(ws) == [(300, dt.date(2026, 6, 30)), (400, dt.date(2026, 7, 31))]


def test_un_corte_futuro_no_gana_sobre_el_ultimo_corte_real():
    """Un 'diciembre' que repite el último valor de un semestral (fecha futura) no desplaza al real."""
    ws = _hoja([(341, dt.date(2026, 6, 30)), (341, dt.date(2026, 12, 31))])  # estado heredado
    nuevos = [
        {"Id": 341, "fecha": dt.date(2026, 12, 31)},  # placeholder futuro
        {"Id": 341, "fecha": dt.date(2026, 6, 30)},  # corte real
        {"Id": 600, "fecha": dt.date(2026, 12, 31)},  # solo hay futuro: se escribe
    ]
    a_escribir, reemplazadas, descartados = reemplazar_cierres_parciales(ws, nuevos, CURSO, HOY)
    assert [(r["Id"], r["fecha"]) for r in a_escribir] == [
        (341, dt.date(2026, 6, 30)),
        (600, dt.date(2026, 12, 31)),
    ]
    assert reemplazadas == 2  # se borran el junio y el diciembre heredados y se reescribe junio
    assert descartados == 1
    assert _contenido(ws) == []


def test_un_corte_futuro_nuevo_no_pisa_un_corte_real_existente():
    ws = _hoja([(341, dt.date(2026, 6, 30))])
    nuevos = [{"Id": 341, "fecha": dt.date(2026, 12, 31)}]
    a_escribir, reemplazadas, descartados = reemplazar_cierres_parciales(ws, nuevos, CURSO, HOY)
    assert a_escribir == [] and (reemplazadas, descartados) == (0, 1)
    assert _contenido(ws) == [(341, dt.date(2026, 6, 30))]


def test_un_año_cerrado_no_se_reemplaza():
    ws = _hoja([(217, dt.date(2025, 12, 31))])
    nuevos = [{"Id": 217, "fecha": dt.date(2025, 12, 31)}]
    a_escribir, reemplazadas, descartados = reemplazar_cierres_parciales(ws, nuevos, CURSO, HOY)
    assert a_escribir == nuevos and (reemplazadas, descartados) == (0, 0)
    assert len(_contenido(ws)) == 1


def test_marcar_estado_cierre_crea_la_columna_y_rotula_cada_fila():
    ws = _hoja([(1, dt.date(2024, 12, 31)), (1, dt.date(2026, 7, 31))])
    assert marcar_estado_cierre(ws, CURSO) == 2
    col = next(c.column for c in ws[1] if c.value == "Estado_Cierre")
    assert [ws.cell(r, col).value for r in (2, 3)] == ["Final", "Parcial"]
    marcar_estado_cierre(ws, CURSO)  # idempotente: no duplica la columna
    assert sum(1 for c in ws[1] if c.value == "Estado_Cierre") == 1


def test_limpiar_cierres_deja_un_cierre_por_año_cerrado_y_todos_los_cortes_del_en_curso(monkeypatch):
    import scripts.etl.purga as purga

    monkeypatch.setattr(purga, "anio_en_curso", lambda: CURSO)
    ws = _hoja(
        [
            (217, dt.date(2025, 6, 30)),
            (217, dt.date(2025, 12, 31)),
            (217, dt.date(2026, 3, 31)),
            (217, dt.date(2026, 6, 30)),
        ]
    )
    limpiar_cierres_existentes(ws)
    assert _contenido(ws) == [
        (217, dt.date(2025, 12, 31)),
        (217, dt.date(2026, 3, 31)),
        (217, dt.date(2026, 6, 30)),
    ]


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        ("Cierre", "Cierre"),
        (" promedio ", "promedio"),
        ("Acumulado", "Acumulado"),
        ("Aplicar la fórmula a cada serie y luego promediar los resultados", ""),
        (None, ""),
    ],
)
def test_tipo_calculo_solo_acepta_cierre_promedio_acumulado(valor, esperado):
    assert _tipo_calculo_valido(valor) == esperado


def test_hoja_catalogo_un_vacio_no_borra_el_tipo_ya_cargado():
    import pandas as pd

    resultado = {
        "extraccion_map": {},
        "tipo_calculo_map": {"7": "Acumulado"},
        "tipo_indicador_map": {},
        "variables_campo_map": {},
        "user_data": {"7": {"TipoCalculo": "Acumulado", "Asociacion": "", "Formato_Valores": ""}},
    }
    _leer_hoja_catalogo(pd.DataFrame({"Id": ["7", "8"], "TipoCalculo": [None, "Promedio"]}), resultado)
    assert resultado["tipo_calculo_map"] == {"7": "Acumulado", "8": "Promedio"}


def test_resalta_sin_borrar_los_cierres_con_id_repetido_en_el_mismo_año():
    ws = _hoja(
        [
            (109, dt.date(2023, 10, 31)),  # mismo Id y año con dos cierres: se resaltan
            (109, dt.date(2023, 12, 31)),
            (109, dt.date(2024, 12, 31)),  # otro año: no
            (7, dt.date(2023, 12, 31)),
        ]
    )
    assert resaltar_duplicados_id_anio(ws) == 2
    assert ws.max_row == 5  # no se borra nada
    color = lambda r: ws.cell(r, 1).fill.fgColor.rgb  # noqa: E731
    assert color(2).endswith("F8CBAD") and color(3).endswith("F8CBAD")
    assert not str(color(4)).endswith("F8CBAD") and not str(color(5)).endswith("F8CBAD")
    # idempotente: si el duplicado desaparece, el resalte también
    ws.delete_rows(3)
    assert resaltar_duplicados_id_anio(ws) == 0
    assert not str(color(2)).endswith("F8CBAD")
